from __future__ import annotations

import hashlib
import importlib
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request
from contextlib import contextmanager
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Dict, Iterator, List
from urllib.parse import urlparse

from app.shared.utils.atomic_file import atomic_write_bytes, validate_image_file
from app.shared.utils.runtime_temp_workspace import resolve_runtime_state_root
from handover_log_module.service.handover_capacity_report_service import HandoverCapacityReportService
from handover_log_module.service.handover_summary_message_service import HandoverSummaryMessageService
from handover_log_module.service.review_link_delivery_service import ReviewLinkDeliveryService
from handover_log_module.service.review_session_service import ReviewSessionNotFoundError, ReviewSessionService
from pipeline_utils import get_app_dir, get_app_root_dir


_SCREENSHOT_QUEUE_TIMEOUT_SEC = 600
_LIBREOFFICE_CONVERT_TIMEOUT_SEC = 180
_DEPENDENCY_INSTALL_TIMEOUT_SEC = 900
_LIBREOFFICE_DOWNLOAD_READ_TIMEOUT_SEC = 120
_PDF_RENDER_SCALE = 2.0
_CAPACITY_REPORT_SHEET_NAME = "本班组"
_PYPI_MIRRORS: tuple[tuple[str, str, str], ...] = (
    ("清华源", "https://pypi.tuna.tsinghua.edu.cn/simple", "pypi.tuna.tsinghua.edu.cn"),
    ("阿里云源", "https://mirrors.aliyun.com/pypi/simple", "mirrors.aliyun.com"),
)
_LIBREOFFICE_INSTALLER_URL_ENV = "LIBREOFFICE_INSTALLER_URLS"
_LIBREOFFICE_ONLINE_INSTALLERS: tuple[tuple[str, str], ...] = (
    (
        "清华源 LibreOffice 26.2.3",
        "https://mirrors.tuna.tsinghua.edu.cn/tdf/libreoffice/stable/26.2.3/win/x86_64/LibreOffice_26.2.3_Win_x86-64.msi",
    ),
    (
        "中科大源 LibreOffice 26.2.3",
        "https://mirrors.ustc.edu.cn/tdf/libreoffice/stable/26.2.3/win/x86_64/LibreOffice_26.2.3_Win_x86-64.msi",
    ),
    (
        "南京大学源 LibreOffice 26.2.3",
        "https://mirrors.nju.edu.cn/tdf/libreoffice/stable/26.2.3/win/x86_64/LibreOffice_26.2.3_Win_x86-64.msi",
    ),
    (
        "官方源 LibreOffice 26.2.3",
        "https://download.documentfoundation.org/libreoffice/stable/26.2.3/win/x86_64/LibreOffice_26.2.3_Win_x86-64.msi",
    ),
    (
        "清华源 LibreOffice 25.8.6",
        "https://mirrors.tuna.tsinghua.edu.cn/tdf/libreoffice/stable/25.8.6/win/x86_64/LibreOffice_25.8.6_Win_x86-64.msi",
    ),
    (
        "官方源 LibreOffice 25.8.6",
        "https://download.documentfoundation.org/libreoffice/stable/25.8.6/win/x86_64/LibreOffice_25.8.6_Win_x86-64.msi",
    ),
)


def _now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _safe_name(value: str) -> str:
    text = _text(value)
    text = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", text, flags=re.UNICODE)
    return text.strip("._") or "unknown"


def _runtime_root(handover_cfg: Dict[str, Any]) -> Path:
    return resolve_runtime_state_root(
        runtime_config={"paths": handover_cfg.get("_global_paths", {}) if isinstance(handover_cfg, dict) else {}},
        app_dir=get_app_dir(),
    )


def _subprocess_creation_flags() -> int:
    if os.name != "nt":
        return 0
    return int(getattr(subprocess, "CREATE_NO_WINDOW", 0) or 0)


def _tail_process_output(value: str, *, limit: int = 1200) -> str:
    text = str(value or "").strip()
    if len(text) <= limit:
        return text
    return text[-limit:]


def _dependency_search_roots() -> List[Path]:
    roots: List[Path] = []
    for candidate in (get_app_dir(), get_app_root_dir(get_app_dir()), Path.cwd()):
        try:
            path = Path(candidate).resolve()
        except Exception:
            path = Path(candidate)
        if path not in roots:
            roots.append(path)
    return roots


def _find_soffice_executable() -> Path | None:
    env_path = _text(os.environ.get("LIBREOFFICE_SOFFICE_PATH"))
    env_home = _text(os.environ.get("LIBREOFFICE_HOME"))
    candidates: List[str] = []
    if env_path:
        candidates.append(env_path)
    if env_home:
        home_path = Path(env_home)
        candidates.append(str(home_path / "program" / "soffice.exe"))
        candidates.append(str(home_path / "soffice.exe"))
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            candidates.append(found)
    for root in _dependency_search_roots():
        candidates.extend(
            [
                str(root / "runtime_dependencies" / "LibreOffice" / "program" / "soffice.exe"),
                str(root / "runtime_dependencies" / "libreoffice" / "program" / "soffice.exe"),
                str(root / "dependencies" / "LibreOffice" / "program" / "soffice.exe"),
                str(root / "tools" / "LibreOffice" / "program" / "soffice.exe"),
                str(root / "LibreOffice" / "program" / "soffice.exe"),
            ]
        )
    candidates.extend(
        [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
    )
    for candidate in candidates:
        path = Path(candidate)
        if path.exists() and path.is_file():
            return path
    return None


def _pypdfium2_importable() -> bool:
    try:
        importlib.import_module("pypdfium2")
        return True
    except Exception:
        return False


def _pypdfium2_install_commands() -> List[tuple[str, List[str]]]:
    base = [
        sys.executable,
        "-m",
        "pip",
        "install",
        "pypdfium2",
        "--timeout",
        "300",
        "--retries",
        "5",
    ]
    output: List[tuple[str, List[str]]] = []
    for label, index_url, trusted_host in _PYPI_MIRRORS:
        output.append((label, [*base, "-i", index_url, "--trusted-host", trusted_host]))
    output.append(("默认源", base))
    return output


def _install_pypdfium2(*, emit_log: Callable[[str], None]) -> bool:
    emit_log("[交接班][容量表图片发送] pypdfium2 未安装，开始自动安装")
    last_error = ""
    for source_label, command in _pypdfium2_install_commands():
        emit_log(f"[交接班][容量表图片发送] pypdfium2 自动安装尝试 source={source_label}")
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=_DEPENDENCY_INSTALL_TIMEOUT_SEC,
                creationflags=_subprocess_creation_flags(),
            )
        except Exception as exc:  # noqa: BLE001
            last_error = str(exc)
            emit_log(f"[交接班][容量表图片发送] pypdfium2 自动安装尝试失败 source={source_label}, error={exc}")
            continue
        if result.returncode != 0:
            last_error = _tail_process_output(result.stderr or result.stdout)
            emit_log(
                "[交接班][容量表图片发送] pypdfium2 自动安装尝试失败 "
                f"source={source_label}, code={result.returncode}, stderr={last_error}"
            )
            continue
        importlib.invalidate_caches()
        if _pypdfium2_importable():
            emit_log(f"[交接班][容量表图片发送] pypdfium2 自动安装完成 source={source_label}")
            return True
        last_error = "pip安装完成但当前Python仍无法导入pypdfium2"
    emit_log(f"[交接班][容量表图片发送] pypdfium2 自动安装失败: {last_error or 'unknown'}")
    return False


def _run_libreoffice_installer(
    installer: Path,
    *,
    emit_log: Callable[[str], None],
    source_label: str,
) -> bool:
    suffix = installer.suffix.lower()
    if suffix == ".msi":
        commands = [["msiexec.exe", "/i", str(installer), "/qn", "/norestart"]]
    elif suffix == ".exe":
        commands = [[str(installer), "/S"], [str(installer), "/quiet", "/norestart"]]
    else:
        emit_log(f"[交接班][容量表图片发送] LibreOffice 安装包格式不支持 source={source_label}, installer={installer}")
        return False
    last_error = ""
    for command in commands:
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=_DEPENDENCY_INSTALL_TIMEOUT_SEC,
                creationflags=_subprocess_creation_flags(),
            )
        except Exception as exc:  # noqa: BLE001
            last_error = str(exc)
            continue
        if result.returncode == 0 and _find_soffice_executable() is not None:
            emit_log(f"[交接班][容量表图片发送] LibreOffice 安装完成 source={source_label}")
            return True
        last_error = _tail_process_output(result.stderr or result.stdout)
    emit_log(
        "[交接班][容量表图片发送] LibreOffice 安装失败 "
        f"source={source_label}, installer={installer}, error={last_error or 'unknown'}"
    )
    return False


def _dependency_download_dir() -> Path:
    env_dir = _text(os.environ.get("QJPT_DEPENDENCY_DOWNLOAD_DIR"))
    if env_dir:
        return Path(env_dir)
    return _dependency_search_roots()[0] / ".runtime" / "dependency_downloads"


def _libreoffice_online_installers() -> List[tuple[str, str]]:
    output: List[tuple[str, str]] = []
    raw_env = _text(os.environ.get(_LIBREOFFICE_INSTALLER_URL_ENV))
    if raw_env:
        for index, item in enumerate(re.split(r"[\r\n;,]+", raw_env), start=1):
            url = _text(item)
            if url:
                output.append((f"环境变量{index}", url))
    seen: set[str] = set()
    deduped: List[tuple[str, str]] = []
    for label, url in [*output, *_LIBREOFFICE_ONLINE_INSTALLERS]:
        cleaned = _text(url)
        if not cleaned or cleaned in seen:
            continue
        seen.add(cleaned)
        deduped.append((label, cleaned))
    return deduped


def _libreoffice_installer_filename(url: str) -> str:
    parsed_name = Path(urlparse(url).path).name
    if parsed_name.lower().endswith((".msi", ".exe")):
        return parsed_name
    digest = hashlib.sha256(url.encode("utf-8", errors="ignore")).hexdigest()[:12]
    return f"LibreOffice_online_{digest}.msi"


def _download_libreoffice_installer(
    url: str,
    *,
    source_label: str,
    emit_log: Callable[[str], None],
) -> Path:
    download_dir = _dependency_download_dir()
    download_dir.mkdir(parents=True, exist_ok=True)
    target = download_dir / _libreoffice_installer_filename(url)
    if target.exists() and target.stat().st_size >= 100 * 1024 * 1024:
        emit_log(
            "[交接班][容量表图片发送] LibreOffice 在线安装包已存在，复用缓存 "
            f"source={source_label}, path={target}, size_mb={target.stat().st_size // 1024 // 1024}"
        )
        return target
    temp_path = target.with_name(target.name + ".part")
    if temp_path.exists():
        try:
            temp_path.unlink()
        except Exception:  # noqa: BLE001
            pass
    emit_log(
        "[交接班][容量表图片发送] LibreOffice 未安装，开始在线下载安装包 "
        f"source={source_label}, url={url}"
    )
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 QJPT-AutoInstaller/1.0"},
    )
    downloaded = 0
    total = 0
    last_log_at = 0
    try:
        with urllib.request.urlopen(request, timeout=_LIBREOFFICE_DOWNLOAD_READ_TIMEOUT_SEC) as response:
            total_text = response.headers.get("Content-Length", "")
            if str(total_text).isdigit():
                total = int(total_text)
            with temp_path.open("wb") as handle:
                while True:
                    chunk = response.read(1024 * 1024)
                    if not chunk:
                        break
                    handle.write(chunk)
                    downloaded += len(chunk)
                    if downloaded - last_log_at >= 50 * 1024 * 1024:
                        last_log_at = downloaded
                        if total > 0:
                            emit_log(
                                "[交接班][容量表图片发送] LibreOffice 安装包下载中 "
                                f"source={source_label}, downloaded_mb={downloaded // 1024 // 1024}, "
                                f"total_mb={total // 1024 // 1024}"
                            )
                        else:
                            emit_log(
                                "[交接班][容量表图片发送] LibreOffice 安装包下载中 "
                                f"source={source_label}, downloaded_mb={downloaded // 1024 // 1024}"
                            )
    except urllib.error.URLError as exc:
        raise RuntimeError(f"下载安装包失败: {exc}") from exc
    except Exception:
        raise
    if downloaded < 100 * 1024 * 1024:
        try:
            temp_path.unlink()
        except Exception:  # noqa: BLE001
            pass
        raise RuntimeError(f"下载安装包大小异常: {downloaded} bytes")
    os.replace(temp_path, target)
    emit_log(
        "[交接班][容量表图片发送] LibreOffice 安装包下载完成 "
        f"source={source_label}, path={target}, size_mb={downloaded // 1024 // 1024}"
    )
    return target


def _install_libreoffice_from_online_download(*, emit_log: Callable[[str], None]) -> bool:
    last_error = ""
    for source_label, url in _libreoffice_online_installers():
        try:
            installer = _download_libreoffice_installer(url, source_label=source_label, emit_log=emit_log)
        except Exception as exc:  # noqa: BLE001
            last_error = str(exc)
            emit_log(
                "[交接班][容量表图片发送] LibreOffice 在线安装包下载失败 "
                f"source={source_label}, error={exc}"
            )
            continue
        if _run_libreoffice_installer(installer, emit_log=emit_log, source_label=source_label):
            return True
        last_error = f"{source_label} 安装失败"
    emit_log(f"[交接班][容量表图片发送] LibreOffice 在线下载安装失败: {last_error or 'unknown'}")
    return False


def _install_libreoffice_windows(*, emit_log: Callable[[str], None]) -> bool:
    if os.name != "nt":
        emit_log("[交接班][容量表图片发送] 非 Windows 环境，无法自动安装 LibreOffice")
        return False
    if _install_libreoffice_from_online_download(emit_log=emit_log):
        return True
    winget = shutil.which("winget")
    if winget:
        base_command = [
            winget,
            "install",
            "--id",
            "TheDocumentFoundation.LibreOffice",
            "-e",
            "--accept-package-agreements",
            "--accept-source-agreements",
            "--silent",
        ]
        commands = [base_command + ["--disable-interactivity"], base_command]
        emit_log("[交接班][容量表图片发送] LibreOffice 在线下载失败，开始通过 winget 自动安装")
        last_error = ""
        for command in commands:
            try:
                result = subprocess.run(
                    command,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=_DEPENDENCY_INSTALL_TIMEOUT_SEC,
                    creationflags=_subprocess_creation_flags(),
                )
            except Exception as exc:  # noqa: BLE001
                last_error = str(exc)
                continue
            if result.returncode == 0 and _find_soffice_executable() is not None:
                emit_log("[交接班][容量表图片发送] LibreOffice winget 自动安装完成")
                return True
            last_error = _tail_process_output(result.stderr or result.stdout)
        emit_log(f"[交接班][容量表图片发送] LibreOffice winget 自动安装失败: {last_error or 'unknown'}")
    else:
        emit_log("[交接班][容量表图片发送] 未找到 winget，跳过 winget 自动安装兜底")
    emit_log("[交接班][容量表图片发送] LibreOffice 自动安装失败，请检查网络、管理员权限或杀毒软件拦截")
    return False


def ensure_capacity_report_image_runtime_dependencies(
    handover_cfg: Dict[str, Any] | None = None,
    *,
    emit_log: Callable[[str], None] = print,
    install_missing: bool = True,
) -> Dict[str, Any]:
    """Ensure the external-side capacity report screenshot renderer can run."""

    started = time.monotonic()
    result: Dict[str, Any] = {
        "ok": True,
        "soffice_path": "",
        "pypdfium2": False,
        "installed": [],
        "errors": [],
    }
    soffice = _find_soffice_executable()
    if soffice is None and install_missing:
        if _install_libreoffice_windows(emit_log=emit_log):
            result["installed"].append("libreoffice")
            soffice = _find_soffice_executable()
    if soffice is None:
        result["ok"] = False
        result["errors"].append("LibreOffice未安装或自动安装失败，请检查网络、管理员权限或杀毒软件拦截")
    else:
        result["soffice_path"] = str(soffice)

    pypdfium_ready = _pypdfium2_importable()
    if not pypdfium_ready and install_missing:
        if _install_pypdfium2(emit_log=emit_log):
            result["installed"].append("pypdfium2")
            pypdfium_ready = True
    if not pypdfium_ready:
        result["ok"] = False
        result["errors"].append("pypdfium2未安装")
    result["pypdfium2"] = bool(pypdfium_ready)
    elapsed_ms = int((time.monotonic() - started) * 1000)
    if result["ok"]:
        emit_log(
            "[交接班][容量表图片发送] 截图依赖已就绪 "
            f"soffice={result['soffice_path']}, pypdfium2={result['pypdfium2']}, elapsed_ms={elapsed_ms}"
        )
    else:
        emit_log(
            "[交接班][容量表图片发送] 截图依赖未完全就绪 "
            f"errors={'; '.join(result['errors'])}, elapsed_ms={elapsed_ms}"
        )
    return result


class CapacityReportImageRenderQueue:
    def __init__(self) -> None:
        self._condition = threading.Condition()
        self._next_ticket = 0
        self._serving_ticket = 0

    @contextmanager
    def acquire(
        self,
        *,
        timeout_sec: int,
        emit_log: Callable[[str], None],
        context: str,
    ) -> Iterator[Dict[str, int]]:
        started = time.monotonic()
        with self._condition:
            ticket = self._next_ticket
            self._next_ticket += 1
            pending_before = max(0, ticket - self._serving_ticket)
            emit_log(
                "[交接班][容量表图片发送] 进入容量表截图队列 "
                f"{context}, ticket={ticket}, pending_before={pending_before}"
            )
            while ticket != self._serving_ticket:
                remaining = timeout_sec - (time.monotonic() - started)
                if remaining <= 0:
                    raise TimeoutError("容量图片截图繁忙，请稍后重试")
                self._condition.wait(timeout=remaining)
        wait_ms = int((time.monotonic() - started) * 1000)
        emit_log(
            "[交接班][容量表图片发送] 已获取容量表截图队列 "
            f"{context}, ticket={ticket}, wait_ms={wait_ms}"
        )
        try:
            yield {"ticket": ticket, "wait_ms": wait_ms}
        finally:
            with self._condition:
                self._serving_ticket += 1
                self._condition.notify_all()
            emit_log(
                "[交接班][容量表图片发送] 已释放容量表截图队列 "
                f"{context}, ticket={ticket}"
            )


_SCREENSHOT_QUEUE = CapacityReportImageRenderQueue()


class CapacityReportImageDeliveryService:
    def __init__(
        self,
        handover_cfg: Dict[str, Any],
        *,
        config_path: str | Path | None = None,
        review_service: ReviewSessionService | None = None,
        link_service: ReviewLinkDeliveryService | None = None,
        capacity_service: HandoverCapacityReportService | None = None,
        summary_service: HandoverSummaryMessageService | None = None,
    ) -> None:
        self.handover_cfg = handover_cfg if isinstance(handover_cfg, dict) else {}
        self.config_path = Path(config_path) if config_path else None
        self.review_service = review_service or ReviewSessionService(self.handover_cfg)
        self.link_service = link_service or ReviewLinkDeliveryService(self.handover_cfg, config_path=self.config_path)
        self.capacity_service = capacity_service or HandoverCapacityReportService(self.handover_cfg)
        self.summary_service = summary_service or HandoverSummaryMessageService(
            self.handover_cfg,
            config_path=self.config_path,
        )

    def _configured_sheet_name(self) -> str:
        capacity_cfg = self.handover_cfg.get("capacity_report", {})
        template_cfg = capacity_cfg.get("template", {}) if isinstance(capacity_cfg, dict) else {}
        configured = _text(template_cfg.get("sheet_name") if isinstance(template_cfg, dict) else "")
        return configured or _CAPACITY_REPORT_SHEET_NAME

    def _libreoffice_temp_root(self) -> Path:
        root = _runtime_root(self.handover_cfg) / "handover" / "capacity_report_images" / "libreoffice"
        root.mkdir(parents=True, exist_ok=True)
        return root

    def _output_image_path(
        self,
        *,
        session: Dict[str, Any],
        signature: str,
    ) -> Path:
        building = _text(session.get("building"))
        duty_date = _text(session.get("duty_date"))
        duty_shift = _text(session.get("duty_shift")).lower()
        session_id = _text(session.get("session_id"))
        session_hash = hashlib.sha1(session_id.encode("utf-8", errors="ignore")).hexdigest()[:10]
        batch_dir = _runtime_root(self.handover_cfg) / "handover" / "capacity_report_images" / _safe_name(f"{duty_date}--{duty_shift}")
        return batch_dir / f"{_safe_name(building)}_{session_hash}_{signature[:12]}.png"

    @staticmethod
    def _image_stat(path: Path) -> Dict[str, int]:
        stat = path.stat()
        return {
            "image_file_size": int(getattr(stat, "st_size", 0) or 0),
            "image_file_mtime_ns": int(getattr(stat, "st_mtime_ns", 0) or int(getattr(stat, "st_mtime", 0) or 0)),
        }

    @staticmethod
    def _cached_image_valid(delivery: Dict[str, Any], signature: str, *, emit_log: Callable[[str], None]) -> Path | None:
        image_signature = _text(delivery.get("image_signature"))
        image_path_text = _text(delivery.get("image_path"))
        if not signature or image_signature != signature or not image_path_text:
            return None
        path = Path(image_path_text)
        if not path.exists() or not path.is_file():
            return None
        try:
            validate_image_file(path)
            stat = path.stat()
            stored_size = int(delivery.get("image_file_size", 0) or 0)
            stored_mtime = int(delivery.get("image_file_mtime_ns", 0) or 0)
            if stored_size and stored_size != int(getattr(stat, "st_size", 0) or 0):
                return None
            current_mtime = int(getattr(stat, "st_mtime_ns", 0) or int(getattr(stat, "st_mtime", 0) or 0))
            if stored_mtime and stored_mtime != current_mtime:
                return None
            return path
        except Exception as exc:  # noqa: BLE001
            emit_log(f"[交接班][容量表图片发送] 图片缓存校验失败 image={path}, error={exc}")
            return None

    def _file_lock_path(self) -> Path:
        root = _runtime_root(self.handover_cfg) / "handover" / "capacity_report_images"
        root.mkdir(parents=True, exist_ok=True)
        return root / ".capacity_report_image_render.lock"

    @contextmanager
    def _runtime_file_lock(
        self,
        *,
        timeout_sec: int,
        context: str,
        emit_log: Callable[[str], None],
    ) -> Iterator[None]:
        path = self._file_lock_path()
        started = time.monotonic()
        fd: int | None = None
        while fd is None:
            try:
                fd = os.open(str(path), os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.write(fd, f"{os.getpid()}|{context}|{_now_text()}".encode("utf-8", errors="ignore"))
            except FileExistsError:
                try:
                    if time.time() - path.stat().st_mtime > max(timeout_sec, 300):
                        path.unlink(missing_ok=True)
                        continue
                except Exception:
                    pass
                if time.monotonic() - started >= timeout_sec:
                    raise TimeoutError("容量图片截图繁忙，请稍后重试")
                time.sleep(0.2)
        try:
            yield
        finally:
            try:
                if fd is not None:
                    os.close(fd)
            finally:
                try:
                    path.unlink(missing_ok=True)
                except Exception as exc:  # noqa: BLE001
                    emit_log(f"[交接班][容量表图片发送] 容量表截图锁文件清理失败 lock={path}, error={exc}")

    def _render_capacity_report_image(
        self,
        *,
        source_path: Path,
        output_path: Path,
        emit_log: Callable[[str], None],
    ) -> Path:
        context = f"source={source_path}"
        with _SCREENSHOT_QUEUE.acquire(
            timeout_sec=_SCREENSHOT_QUEUE_TIMEOUT_SEC,
            emit_log=emit_log,
            context=context,
        ):
            with self._runtime_file_lock(
                timeout_sec=_SCREENSHOT_QUEUE_TIMEOUT_SEC,
                context=context,
                emit_log=emit_log,
            ):
                return self._render_libreoffice_sheet_image(
                    source_path=source_path,
                    output_path=output_path,
                    emit_log=emit_log,
                )

    @staticmethod
    def _crop_rendered_page(image: Any) -> Any:
        from PIL import Image, ImageChops

        rgb = image.convert("RGB") if getattr(image, "mode", "") != "RGB" else image
        background = Image.new("RGB", rgb.size, "white")
        diff = ImageChops.difference(rgb, background)
        bbox = diff.getbbox()
        if not bbox:
            return rgb
        margin = 8
        left = max(0, bbox[0] - margin)
        upper = max(0, bbox[1] - margin)
        right = min(rgb.width, bbox[2] + margin)
        lower = min(rgb.height, bbox[3] + margin)
        return rgb.crop((left, upper, right, lower))

    def _copy_single_sheet_workbook(
        self,
        *,
        source_path: Path,
        target_path: Path,
        emit_log: Callable[[str], None],
    ) -> str:
        import openpyxl

        configured_sheet = self._configured_sheet_name()
        workbook = openpyxl.load_workbook(source_path)
        try:
            if configured_sheet not in workbook.sheetnames:
                raise RuntimeError(f"未找到截图Sheet: {configured_sheet}")
            target_sheet = workbook[configured_sheet]
            workbook.active = workbook.index(target_sheet)
            for sheet in list(workbook.worksheets):
                if sheet.title != configured_sheet:
                    workbook.remove(sheet)
            target_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(target_path)
            emit_log(
                "[交接班][容量表图片发送] 已生成单Sheet截图临时表 "
                f"sheet={configured_sheet}, temp={target_path}"
            )
            return configured_sheet
        finally:
            workbook.close()

    def _convert_xlsx_to_pdf_with_libreoffice(
        self,
        *,
        source_path: Path,
        output_dir: Path,
        emit_log: Callable[[str], None],
    ) -> Path:
        soffice = _find_soffice_executable()
        if soffice is None:
            raise RuntimeError("LibreOffice未安装，无法生成容量报表图片")
        output_dir.mkdir(parents=True, exist_ok=True)
        profile_dir = self._libreoffice_temp_root() / "profile"
        profile_dir.mkdir(parents=True, exist_ok=True)
        command = [
            str(soffice),
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nodefault",
            "--nolockcheck",
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--convert-to",
            "pdf:calc_pdf_Export",
            "--outdir",
            str(output_dir),
            str(source_path),
        ]
        started = time.monotonic()
        emit_log(
            "[交接班][容量表图片发送] LibreOffice导出PDF开始 "
            f"source={source_path}, outdir={output_dir}, soffice={soffice}"
        )
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=_LIBREOFFICE_CONVERT_TIMEOUT_SEC,
            creationflags=_subprocess_creation_flags(),
        )
        elapsed_ms = int((time.monotonic() - started) * 1000)
        if result.returncode != 0:
            raise RuntimeError(
                "LibreOffice导出PDF失败 "
                f"code={result.returncode}, stderr={_tail_process_output(result.stderr or result.stdout)}"
            )
        pdf_path = output_dir / f"{source_path.stem}.pdf"
        if not pdf_path.exists() or not pdf_path.is_file():
            candidates = sorted(output_dir.glob("*.pdf"), key=lambda item: item.stat().st_mtime, reverse=True)
            pdf_path = candidates[0] if candidates else pdf_path
        if not pdf_path.exists() or not pdf_path.is_file():
            raise RuntimeError("LibreOffice导出PDF成功但未找到PDF文件")
        emit_log(
            "[交接班][容量表图片发送] LibreOffice导出PDF完成 "
            f"pdf={pdf_path}, elapsed_ms={elapsed_ms}"
        )
        return pdf_path

    def _render_pdf_to_image(
        self,
        *,
        pdf_path: Path,
        output_path: Path,
        sheet_name: str,
        emit_log: Callable[[str], None],
    ) -> Path:
        try:
            import pypdfium2
            from PIL import Image
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError(f"pypdfium2/Pillow未安装，无法渲染PDF: {exc}") from exc

        document = pypdfium2.PdfDocument(str(pdf_path))
        pages: List[Any] = []
        try:
            page_count = len(document)
            if page_count <= 0:
                raise RuntimeError("PDF没有可渲染页面")
            for index in range(page_count):
                page = document[index]
                try:
                    bitmap = page.render(scale=_PDF_RENDER_SCALE)
                    page_image = bitmap.to_pil().convert("RGB")
                    pages.append(self._crop_rendered_page(page_image))
                finally:
                    try:
                        page.close()
                    except Exception:
                        pass
        finally:
            try:
                document.close()
            except Exception:
                pass
        if not pages:
            raise RuntimeError("PDF渲染为空")

        gap = 12 if len(pages) > 1 else 0
        width = max(image.width for image in pages)
        height = sum(image.height for image in pages) + gap * max(0, len(pages) - 1)
        combined = Image.new("RGB", (width, height), "white")
        y = 0
        for image in pages:
            x = int((width - image.width) / 2)
            combined.paste(image, (x, y))
            y += image.height + gap

        buffer = BytesIO()
        combined.save(buffer, format="PNG", optimize=True)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        atomic_write_bytes(output_path, buffer.getvalue(), validator=validate_image_file, temp_suffix=".tmp")
        stat = output_path.stat()
        emit_log(
            "[交接班][容量表图片发送] LibreOffice截图生成成功 "
            f"image={output_path}, size={int(getattr(stat, 'st_size', 0) or 0)}, "
            f"sheet={sheet_name}, pdf_pages={len(pages)}, width={width}, height={height}"
        )
        return output_path

    def _render_libreoffice_sheet_image(
        self,
        *,
        source_path: Path,
        output_path: Path,
        emit_log: Callable[[str], None],
    ) -> Path:
        started = time.monotonic()
        temp_root = self._libreoffice_temp_root() / f"job_{os.getpid()}_{int(time.time() * 1000)}"
        try:
            work_dir = temp_root / "work"
            out_dir = temp_root / "pdf"
            sheet_xlsx = work_dir / source_path.name
            sheet_name = self._copy_single_sheet_workbook(
                source_path=source_path,
                target_path=sheet_xlsx,
                emit_log=emit_log,
            )
            pdf_path = self._convert_xlsx_to_pdf_with_libreoffice(
                source_path=sheet_xlsx,
                output_dir=out_dir,
                emit_log=emit_log,
            )
            image_path = self._render_pdf_to_image(
                pdf_path=pdf_path,
                output_path=output_path,
                sheet_name=sheet_name,
                emit_log=emit_log,
            )
            elapsed_ms = int((time.monotonic() - started) * 1000)
            emit_log(
                "[交接班][容量表图片发送] LibreOffice单Sheet截图完成 "
                f"source={source_path}, image={image_path}, elapsed_ms={elapsed_ms}"
            )
            return image_path
        finally:
            try:
                shutil.rmtree(temp_root, ignore_errors=True)
            except Exception as exc:  # noqa: BLE001
                emit_log(f"[交接班][容量表图片发送] LibreOffice截图临时目录清理失败 dir={temp_root}, error={exc}")

    def _persist_delivery(
        self,
        *,
        session_id: str,
        delivery: Dict[str, Any],
        emit_log: Callable[[str], None],
    ) -> Dict[str, Any]:
        try:
            session = self.review_service.update_capacity_image_delivery(
                session_id=session_id,
                capacity_image_delivery=delivery,
            )
            return session.get("capacity_image_delivery", delivery) if isinstance(session, dict) else delivery
        except (ReviewSessionNotFoundError, Exception) as exc:  # noqa: BLE001
            emit_log(f"[交接班][容量表图片发送] 状态更新失败 session_id={session_id}, error={exc}")
            return delivery

    def _persist_review_delivery(
        self,
        *,
        session: Dict[str, Any],
        status: str,
        attempt_at: str,
        successful_recipients: List[str],
        failed_recipients: List[Dict[str, str]],
        error: str,
        emit_log: Callable[[str], None],
    ) -> Dict[str, Any]:
        session_id = _text(session.get("session_id"))
        previous = session.get("review_link_delivery", {}) if isinstance(session.get("review_link_delivery", {}), dict) else {}
        payload = {
            **previous,
            "status": status,
            "last_attempt_at": attempt_at,
            "last_sent_at": attempt_at if status == "success" else _text(previous.get("last_sent_at")),
            "error": error,
            "successful_recipients": successful_recipients,
            "failed_recipients": failed_recipients,
            "source": "capacity_image_send",
        }
        try:
            updated = self.review_service.update_review_link_delivery(
                session_id=session_id,
                review_link_delivery=payload,
            )
            return updated.get("review_link_delivery", payload) if isinstance(updated, dict) else payload
        except Exception as exc:  # noqa: BLE001
            emit_log(f"[交接班][容量表图片发送] 审核文本发送状态更新失败 session_id={session_id}, error={exc}")
            return payload

    def _fail(
        self,
        *,
        session: Dict[str, Any],
        delivery: Dict[str, Any],
        attempt_at: str,
        error: str,
        failed_recipients: List[Dict[str, str]] | None = None,
        emit_log: Callable[[str], None],
        update_review_delivery: bool = False,
    ) -> Dict[str, Any]:
        session_id = _text(session.get("session_id"))
        building = _text(session.get("building"))
        failed = list(failed_recipients or [])
        next_delivery = {
            **delivery,
            "status": "failed",
            "last_attempt_at": attempt_at,
            "error": error,
            "failed_recipients": failed,
            "source": "manual",
        }
        persisted = self._persist_delivery(session_id=session_id, delivery=next_delivery, emit_log=emit_log)
        if update_review_delivery:
            review_delivery = self._persist_review_delivery(
                session=session,
                status="failed",
                attempt_at=attempt_at,
                successful_recipients=[],
                failed_recipients=failed,
                error=error,
                emit_log=emit_log,
            )
        else:
            review_delivery = (
                dict(session.get("review_link_delivery", {}))
                if isinstance(session.get("review_link_delivery", {}), dict)
                else {}
            )
        emit_log(
            "[交接班][容量表图片发送] 失败 "
            f"building={building}, session_id={session_id}, error={error}, failed_recipients={failed}"
        )
        return {
            "ok": False,
            "status": "failed",
            "error": error,
            "building": building,
            "session_id": session_id,
            "successful_recipients": [],
            "failed_recipients": failed,
            "capacity_image_delivery": persisted,
            "review_link_delivery": review_delivery,
        }

    def send_for_session(
        self,
        session: Dict[str, Any],
        *,
        building: str = "",
        handover_cells: Dict[str, Any] | None = None,
        shared_110kv: Dict[str, Any] | None = None,
        cooling_pump_pressures: Dict[str, Any] | None = None,
        client_id: str = "",
        ensure_capacity_ready: Callable[[], Dict[str, Any]] | None = None,
        emit_log: Callable[[str], None] = print,
    ) -> Dict[str, Any]:
        current_session = dict(session if isinstance(session, dict) else {})
        session_id = _text(current_session.get("session_id"))
        building_text = _text(building) or _text(current_session.get("building"))
        current_session["building"] = building_text
        attempt_at = _now_text()
        delivery_state = (
            dict(current_session.get("capacity_image_delivery", {}))
            if isinstance(current_session.get("capacity_image_delivery", {}), dict)
            else {}
        )
        delivery_state = {
            **delivery_state,
            "status": "sending",
            "last_attempt_at": attempt_at,
            "error": "",
            "source": "manual",
            "failed_recipients": [],
        }
        self._persist_delivery(session_id=session_id, delivery=delivery_state, emit_log=emit_log)

        capacity_file = _text(current_session.get("capacity_output_file"))
        emit_log(
            "[交接班][容量表图片发送] 开始发送 "
            f"building={building_text}, session={session_id}, source={capacity_file}, client_id={_text(client_id) or '-'}"
        )
        if not session_id or not building_text:
            return self._fail(
                session=current_session,
                delivery=delivery_state,
                attempt_at=attempt_at,
                error="容量图片发送缺少 session_id/building",
                emit_log=emit_log,
            )
        capacity_path = Path(capacity_file)
        if not capacity_file or not capacity_path.exists() or not capacity_path.is_file():
            return self._fail(
                session=current_session,
                delivery=delivery_state,
                attempt_at=attempt_at,
                error="交接班容量报表文件不存在",
                emit_log=emit_log,
            )

        recipient_snapshot = self.link_service._recipient_snapshot_for_building(building_text)
        recipients = list(recipient_snapshot.get("recipients", []))
        emit_log(
            "[交接班][容量表图片发送] 收件人读取完成 "
            f"building={building_text}, session={session_id}, recipients={len(recipients)}, "
            f"raw={int(recipient_snapshot.get('raw_count', 0) or 0)}, enabled={int(recipient_snapshot.get('enabled_count', 0) or 0)}, "
            f"disabled={int(recipient_snapshot.get('disabled_count', 0) or 0)}, invalid={int(recipient_snapshot.get('invalid_count', 0) or 0)}, "
            f"open_ids={recipient_snapshot.get('open_ids', [])}"
        )
        if not recipients:
            return self._fail(
                session=current_session,
                delivery=delivery_state,
                attempt_at=attempt_at,
                error="当前楼未配置启用且有效的审核接收人 open_id",
                emit_log=emit_log,
            )

        message_text = self.summary_service.build_for_session(current_session, emit_log=emit_log)
        if not message_text:
            return self._fail(
                session=current_session,
                delivery=delivery_state,
                attempt_at=attempt_at,
                error="交接班日志全文生成失败，未发送",
                emit_log=emit_log,
            )
        emit_log(
            "[交接班][容量表图片发送] 审核文本生成完成 "
            f"building={building_text}, session={session_id}, length={len(message_text)}"
        )

        cells = handover_cells if isinstance(handover_cells, dict) else {}
        signature_info = self.capacity_service.build_capacity_overlay_signature(
            building=building_text,
            duty_date=_text(current_session.get("duty_date")),
            duty_shift=_text(current_session.get("duty_shift")).lower(),
            handover_cells=cells,
            capacity_output_file=capacity_file,
            shared_110kv=shared_110kv,
            cooling_pump_pressures=cooling_pump_pressures,
            client_id=client_id,
            emit_log=emit_log,
        )
        target_signature = _text(signature_info.get("signature"))
        sync_state = current_session.get("capacity_sync", {}) if isinstance(current_session.get("capacity_sync", {}), dict) else {}
        sync_ready = (
            _text(sync_state.get("status")).lower() == "ready"
            and _text(sync_state.get("input_signature")) == _text(signature_info.get("input_signature"))
        )
        cached_image = self._cached_image_valid(delivery_state, target_signature, emit_log=emit_log) if sync_ready else None
        emit_log(
            "[交接班][容量表图片发送] 容量图片签名检查 "
            f"building={building_text}, session={session_id}, signature={target_signature[:12]}, "
            f"capacity_sync_ready={bool(sync_ready)}, cache_hit={bool(cached_image)}"
        )

        if cached_image is None:
            if callable(ensure_capacity_ready):
                try:
                    updated_session = ensure_capacity_ready()
                    if isinstance(updated_session, dict):
                        current_session.update(updated_session)
                        current_session["building"] = building_text
                except Exception as exc:  # noqa: BLE001
                    return self._fail(
                        session=current_session,
                        delivery=delivery_state,
                        attempt_at=attempt_at,
                        error=f"容量报表补写失败: {exc}",
                        emit_log=emit_log,
                    )
            else:
                sync_payload = self.capacity_service.sync_overlay_for_existing_report_from_cells(
                    building=building_text,
                    duty_date=_text(current_session.get("duty_date")),
                    duty_shift=_text(current_session.get("duty_shift")).lower(),
                    handover_cells=cells,
                    capacity_output_file=capacity_file,
                    shared_110kv=shared_110kv,
                    cooling_pump_pressures=cooling_pump_pressures,
                    client_id=client_id,
                    emit_log=emit_log,
                )
                sync_status = _text(sync_payload.get("status")).lower() if isinstance(sync_payload, dict) else "failed"
                try:
                    current_session.update(
                        self.review_service.update_capacity_sync(
                            session_id=session_id,
                            capacity_sync=sync_payload,
                            capacity_status="success" if sync_status == "ready" else sync_status,
                            capacity_error="" if sync_status == "ready" else _text(sync_payload.get("error")),
                        )
                    )
                except Exception as exc:  # noqa: BLE001
                    emit_log(f"[交接班][容量表图片发送] 容量补写状态更新失败 session={session_id}, error={exc}")

            sync_state = current_session.get("capacity_sync", {}) if isinstance(current_session.get("capacity_sync", {}), dict) else {}
            if _text(sync_state.get("status")).lower() != "ready":
                return self._fail(
                    session=current_session,
                    delivery=delivery_state,
                    attempt_at=attempt_at,
                    error=_text(sync_state.get("error")) or "容量报表补写未就绪，未发送",
                    emit_log=emit_log,
                )
            signature_info = self.capacity_service.build_capacity_overlay_signature(
                building=building_text,
                duty_date=_text(current_session.get("duty_date")),
                duty_shift=_text(current_session.get("duty_shift")).lower(),
                handover_cells=cells,
                capacity_output_file=capacity_file,
                shared_110kv=shared_110kv,
                cooling_pump_pressures=cooling_pump_pressures,
                client_id=client_id,
                emit_log=emit_log,
            )
            target_signature = _text(signature_info.get("signature"))
            if not bool(signature_info.get("valid", False)):
                return self._fail(
                    session=current_session,
                    delivery=delivery_state,
                    attempt_at=attempt_at,
                    error=_text(signature_info.get("error")) or "容量报表补写输入不完整，未发送",
                    emit_log=emit_log,
                )
            output_image = self._output_image_path(session=current_session, signature=target_signature)
            try:
                image_path = self._render_capacity_report_image(
                    source_path=capacity_path,
                    output_path=output_image,
                    emit_log=emit_log,
                )
            except TimeoutError as exc:
                return self._fail(
                    session=current_session,
                    delivery=delivery_state,
                    attempt_at=attempt_at,
                    error=str(exc),
                    emit_log=emit_log,
                )
            except Exception as exc:  # noqa: BLE001
                return self._fail(
                    session=current_session,
                    delivery=delivery_state,
                    attempt_at=attempt_at,
                    error=f"容量表截图失败，未发送: {exc}",
                    emit_log=emit_log,
                )
        else:
            image_path = cached_image

        image_stat = self._image_stat(image_path)
        delivery_state = {
            **delivery_state,
            "image_path": str(image_path),
            "image_signature": target_signature,
            **image_stat,
        }
        self._persist_delivery(session_id=session_id, delivery=delivery_state, emit_log=emit_log)

        try:
            client = self.link_service._build_feishu_client()
            emit_log(f"[交接班][容量表图片发送] 飞书图片上传开始 image={image_path}")
            upload_result = client.upload_image(str(image_path))
            image_key = _text(upload_result.get("image_key"))
            if not image_key:
                raise RuntimeError("飞书图片上传未返回 image_key")
            emit_log(f"[交接班][容量表图片发送] 飞书图片上传成功 image_key={image_key}")
        except Exception as exc:  # noqa: BLE001
            return self._fail(
                session=current_session,
                delivery=delivery_state,
                attempt_at=attempt_at,
                error=f"飞书图片上传失败: {exc}",
                emit_log=emit_log,
            )

        successful_recipients: List[str] = []
        failed_recipients: List[Dict[str, str]] = []
        for recipient in recipients:
            open_id = _text(recipient.get("open_id"))
            note = _text(recipient.get("note"))
            receive_id_type = self.link_service._resolve_effective_receive_id_type(open_id, "open_id")
            try:
                emit_log(
                    "[交接班][容量表图片发送] 文本发送开始 "
                    f"building={building_text}, session_id={session_id}, open_id={open_id}, receive_id_type={receive_id_type}, note={note or '-'}"
                )
                client.send_text_message(
                    receive_id=open_id,
                    receive_id_type=receive_id_type,
                    text=message_text,
                )
                emit_log(
                    "[交接班][容量表图片发送] 文本发送成功 "
                    f"building={building_text}, session_id={session_id}, open_id={open_id}, receive_id_type={receive_id_type}, note={note or '-'}"
                )
            except Exception as exc:  # noqa: BLE001
                failed_recipients.append({"open_id": open_id, "note": note, "step": "text", "error": str(exc)})
                emit_log(
                    "[交接班][容量表图片发送] 文本发送失败 "
                    f"building={building_text}, session_id={session_id}, open_id={open_id}, receive_id_type={receive_id_type}, note={note or '-'}, error={exc}"
                )
                continue
            try:
                emit_log(
                    "[交接班][容量表图片发送] 图片发送开始 "
                    f"building={building_text}, session_id={session_id}, open_id={open_id}, receive_id_type={receive_id_type}, note={note or '-'}"
                )
                client.send_image_message(
                    receive_id=open_id,
                    receive_id_type=receive_id_type,
                    image_key=image_key,
                )
                successful_recipients.append(open_id)
                emit_log(
                    "[交接班][容量表图片发送] 图片发送成功 "
                    f"building={building_text}, session_id={session_id}, open_id={open_id}, receive_id_type={receive_id_type}, note={note or '-'}"
                )
            except Exception as exc:  # noqa: BLE001
                failed_recipients.append({"open_id": open_id, "note": note, "step": "image", "error": str(exc)})
                emit_log(
                    "[交接班][容量表图片发送] 图片发送失败 "
                    f"building={building_text}, session_id={session_id}, open_id={open_id}, receive_id_type={receive_id_type}, note={note or '-'}, error={exc}"
                )

        status = "success" if not failed_recipients and len(successful_recipients) == len(recipients) else "failed"
        error = "" if status == "success" else "审核文本和容量表图片发送失败，详见收件人明细"
        final_delivery = {
            **delivery_state,
            "status": status,
            "last_attempt_at": attempt_at,
            "last_sent_at": attempt_at if status == "success" else _text(delivery_state.get("last_sent_at")),
            "error": error,
            "image_key": image_key,
            "successful_recipients": successful_recipients,
            "failed_recipients": failed_recipients,
            "source": "manual",
        }
        persisted_delivery = self._persist_delivery(session_id=session_id, delivery=final_delivery, emit_log=emit_log)
        review_delivery = self._persist_review_delivery(
            session=current_session,
            status=status,
            attempt_at=attempt_at,
            successful_recipients=successful_recipients,
            failed_recipients=failed_recipients,
            error=error,
            emit_log=emit_log,
        )
        emit_log(
            "[交接班][容量表图片发送] 完成 "
            f"building={building_text}, session_id={session_id}, status={status}, successful={len(successful_recipients)}, "
            f"failed={len(failed_recipients)}, failed_recipients={failed_recipients}"
        )
        return {
            "ok": status == "success",
            "status": status,
            "error": error,
            "building": building_text,
            "session_id": session_id,
            "successful_recipients": successful_recipients,
            "failed_recipients": failed_recipients,
            "capacity_image_delivery": persisted_delivery,
            "review_link_delivery": review_delivery,
        }
