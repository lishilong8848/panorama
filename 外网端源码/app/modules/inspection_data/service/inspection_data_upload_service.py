from __future__ import annotations

import copy
import json
import math
import re
import time
from collections import defaultdict
from datetime import datetime, timedelta, timezone, time as dt_time
from pathlib import Path, PurePosixPath
from typing import Any, Callable

from openpyxl import load_workbook

from app.config.inspection_data_config import INSPECTION_DATA_UPLOAD_DEFAULTS as DEFAULT_CONFIG
from app.config.inspection_data_config import INSPECTION_SOURCE_CONTRACTS as SOURCE_FAMILIES
from app.modules.feishu.service.bitable_client_runtime import FeishuBitableClient
from app.modules.feishu.service.feishu_auth_resolver import require_feishu_auth_settings
from app.modules.internal_bridge_http.service.client import InternalBridgeHttpClient
from app.shared.utils.atomic_file import atomic_write_text


CATALOG_PATH = Path(__file__).resolve().parents[1] / "data" / "inspection_catalog.json"
BACKUP_ROOT = Path(__file__).resolve().parents[4] / ".runtime" / "inspection_data" / "backups"
LOCAL_TZ = timezone(timedelta(hours=8))


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _key(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value)).replace("（", "(").replace("）", ")")


def load_catalog() -> list[dict[str, Any]]:
    return json.loads(CATALOG_PATH.read_text(encoding="utf-8"))["devices"]


def field_definitions() -> list[dict[str, Any]]:
    definitions = {
        name: {"field_name": name, "type": kind}
        for name, kind in {
            "设备编号": 1, "楼栋": 3, "专业": 3, "位置": 1, "巡检表单": 1,
            "上传日期": 5, "源文件时间": 5, "源文件": 1, "启停状态": 3,
            "数据完整性": 3, "缺少源数据": 1, "备注": 1,
        }.items()
    }
    options = {"楼栋": ["D楼", "E楼"], "专业": ["暖通", "电气"],
               "启停状态": ["开启", "关闭"], "数据完整性": ["完整", "缺少源数据"]}
    for name, values in options.items():
        definitions[name]["property"] = {"options": [{"name": v} for v in values]}
    for name in ("上传日期", "源文件时间"):
        definitions[name]["property"] = {"date_formatter": "yyyy/MM/dd HH:mm", "auto_fill": False}
    for device in load_catalog():
        for point in device["points"]:
            definitions.setdefault(point["field"], {"field_name": point["field"], "type": point["type"]})
    return list(definitions.values())


def _point_aliases(point: dict[str, Any]) -> set[str]:
    field, label = _key(point["field"]), _key(point["label"])
    phase = field[:2] if field.startswith(("L1", "L2", "L3")) else ""
    # A phase-less temperature must never select one of three phase readings.
    aliases = {field, field.split("(", 1)[0]}
    if not phase:
        aliases.update((label, label.split("(", 1)[0]))
    else:
        aliases.update((phase + label, phase + label.split("(", 1)[0]))
    return aliases


def _numeric(value: Any, field: str) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    text = _text(value)
    unit = field.rsplit("(", 1)[-1].rstrip(")") if "(" in field else ""
    if unit and text.lower().endswith(unit.lower()):
        text = text[:-len(unit)].strip()
    try:
        result = float(text.replace(",", ""))
    except (ValueError, TypeError):
        return None
    return result if math.isfinite(result) else None


def _on_off(value: Any) -> str:
    # Only explicit on/off points use this mapping, never alarm or switch health.
    text = _text(value)
    if text in {"1", "1.0", "开启", "启动", "运行"}:
        return "开启"
    if text in {"0", "0.0", "关闭", "停止", "未运行", "△", "Δ", "▲"}:
        return "关闭"
    return ""


def _device_aliases(device: dict[str, Any]) -> set[str]:
    name = _key(device["device"])
    names = {name}
    # The room and CRAH number are both printed in the form, not inferred from order.
    match = re.fullmatch(r"CRAH(\d+)", name)
    room = re.search(r"\((\d{3})\)", _key(device["position"]))
    if match and room:
        names.add(f"{room[1]}-CRAH-{match[1]}")
    prefix = device["building"][0] + "-"
    return names | {n[len(prefix):] if n.startswith(prefix) else prefix + n for n in names}


def _first_value_column(sheet) -> tuple[int, str]:
    candidates = []
    extra_values = False
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=40, min_col=5, max_col=128, values_only=True), 1):
        dates = []
        for column, value in enumerate(row, 5):
            extra_values = extra_values or (column > 5 and value is not None)
            if isinstance(value, (datetime, dt_time)) or re.fullmatch(
                r"(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}[ T\n]+)?\d{1,2}:\d{2}(?::\d{2})?", _text(value)
            ):
                dates.append((column, _text(value)))
        if dates:
            candidates.append((len(dates), -row_number, dates[0]))
    if candidates:
        return max(candidates, key=lambda item: item[:2])[2]
    if extra_values:
        raise ValueError("巡检源文件有多列数值但没有可识别的时间表头，无法确定第一个时刻；未修改多维表")
    return 5, ""


class InspectionDataUploadService:
    def __init__(self, runtime_config: dict[str, Any], *, internal_client=None, bitable_client=None):
        self.runtime_config = runtime_config
        raw = runtime_config.get("inspection_data_upload", {})
        self.cfg = {**copy.deepcopy(DEFAULT_CONFIG), **raw}
        self.cfg.pop("source_family", None)
        self.cfg["source_families"] = copy.deepcopy(DEFAULT_CONFIG["source_families"])
        self._internal = internal_client
        self._bitable = bitable_client

    @staticmethod
    def normalize_source_date(value: Any = None) -> str:
        if not value:
            return datetime.now().strftime("%Y-%m-%d")
        return datetime.strptime(str(value), "%Y-%m-%d").strftime("%Y-%m-%d")

    def make_bitable_client(self):
        if self._bitable is not None:
            return self._bitable
        auth = require_feishu_auth_settings(self.runtime_config)
        return FeishuBitableClient(
            **auth, app_token=self.cfg["app_token"], calc_table_id="", attachment_table_id="",
            date_text_to_timestamp_ms_fn=lambda **_: 0,
            canonical_metric_name_fn=lambda value: str(value), dimension_mapping={},
        )

    def ensure_schema(self, *, apply: bool = False) -> dict[str, Any]:
        client = self.make_bitable_client()
        report = {}
        for building in self.cfg["buildings"]:
            table = self.cfg["tables"][building]
            existing = {item["field_name"]: item for item in client.list_fields(table)}
            conflicts = [d["field_name"] for d in field_definitions()
                         if d["field_name"] in existing and existing[d["field_name"]]["type"] != d["type"]]
            if conflicts:
                raise ValueError(f"{building} 巡检字段类型不符，未修改已有字段: {','.join(conflicts)}")
            missing = [d for d in field_definitions() if d["field_name"] not in existing]
            if apply:
                for definition in missing:
                    url = client.LIST_FIELD_URL.format(app_token=client.app_token, table_id=table)
                    client._post_json(url, definition)
                actual = {d["field_name"]: d["type"] for d in client.list_fields(table)}
                if any(actual.get(d["field_name"]) != d["type"] for d in field_definitions()):
                    raise RuntimeError(f"{building} 巡检字段创建后复核失败")
            report[building] = {"table_id": table, "missing_fields": [d["field_name"] for d in missing], "applied": apply}
        return report

    @staticmethod
    def parse_workbook(file_path: Path | str, *, building: str, catalog=None) -> dict[str, Any]:
        devices = [d for d in (load_catalog() if catalog is None else catalog) if d["building"] == building]
        if not devices:
            raise ValueError(f"不支持的巡检楼栋: {building}")
        by_device = defaultdict(list)
        for i, device in enumerate(devices):
            for name in _device_aliases(device):
                by_device[name].append(i)
        readings: dict[int, dict[str, Any]] = defaultdict(dict)
        diagnostics: list[str] = []
        valid_count = 0
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            value_column, value_time = _first_value_column(sheet)
            position = device_name = ""
            for row_number, values in enumerate(sheet.iter_rows(min_col=2, max_col=value_column, values_only=True), 1):
                if row_number > 100000:
                    raise ValueError("巡检源文件超过100000行，停止解析，未修改多维表")
                b, c, d = values[:3]
                raw = values[-1]
                explicit = set(re.findall(r"(?<![A-Za-z0-9])([A-E])(?:楼|栋|-\d{3})", _text(b) + " " + _text(c)))
                if explicit and explicit != {building[0]}:
                    raise ValueError(f"{building} 巡检源文件出现其他楼栋标识，第{row_number}行，未修改多维表")
                if _text(b):
                    if _key(b) != _key(position):
                        device_name = ""
                    position = _text(b)
                if _text(c):
                    device_name = _text(c)
                if not _text(d) or not device_name:
                    continue
                if _key(c) in {"设备", "设备编号", "监控对象"} and _key(d) in {"点位", "测点", "监控指标", "指标"}:
                    continue
                candidates = list(set(by_device.get(_key(device_name), [])))
                if len(candidates) > 1:
                    candidates = [i for i in candidates if _key(devices[i]["position"]) == _key(position)]
                if len(candidates) != 1:
                    diagnostics.append(f"第{row_number}行设备未唯一匹配: {position}/{device_name}/{_text(d)}")
                    continue
                i = candidates[0]
                metric = _key(d)
                if metric == "启停状态":
                    field, value = "启停状态", _on_off(raw)
                else:
                    points = [p for p in devices[i]["points"] if metric in _point_aliases(p)]
                    if len(points) != 1:
                        diagnostics.append(f"第{row_number}行点位未唯一匹配: {device_name}/{_text(d)}")
                        continue
                    p = points[0]
                    field = p["field"]
                    value = _numeric(raw, field) if p["type"] == 2 else _text(raw)
                if value is None or value == "":
                    diagnostics.append(f"第{row_number}行空值或无效数值: {device_name}/{field}")
                    continue
                if field in readings[i] and readings[i][field] != value:
                    raise ValueError(f"{building} 巡检源文件重复点位数值冲突: {device_name}/{field}，未修改多维表")
                readings[i][field] = value
                if field == "运行状态" and _on_off(raw):
                    readings[i]["启停状态"] = _on_off(raw)
                valid_count += 1
        finally:
            workbook.close()
        if not valid_count:
            raise ValueError(f"{building} 巡检源文件没有可识别的有效测点，保留目标表原数据")
        rows = []
        for i, device in enumerate(devices):
            values = readings.get(i, {})
            missing = [p["field"] for p in device["points"] if p["field"] not in values]
            if "启停状态" not in values:
                missing.append("启停状态")
            rows.append({"设备编号": device["device"], "楼栋": building, "专业": device["discipline"],
                         "位置": device["position"], "巡检表单": device["sheet"], **values,
                         "数据完整性": "缺少源数据" if missing else "完整", "缺少源数据": "、".join(missing)})
        return {"rows": rows, "valid_points": valid_count, "matched_devices": len(readings),
                "missing_devices": len(devices) - len(readings), "diagnostics": diagnostics,
                "value_column": value_column, "value_time": value_time}

    def _entry_path(self, entry: dict[str, Any]) -> Path | None:
        relative = _text(entry.get("relative_path")).replace("\\", "/")
        bridge = self.runtime_config.get("shared_bridge", {})
        root = bridge.get("external_root_dir") or bridge.get("root_dir")
        if relative and root:
            part = PurePosixPath(relative)
            if part.is_absolute() or ".." in part.parts or ":" in relative:
                return None
            path = Path(root).joinpath(*part.parts)
        else:
            path = Path(_text(entry.get("file_path") or entry.get("resolved_file_path")))
        try:
            return path if path.suffix.lower() == ".xlsx" and path.is_file() and path.stat().st_size > 0 else None
        except OSError:
            return None

    def _sources(self, source_date: str, log, cancel) -> dict[str, dict[str, dict[str, Any]]]:
        internal = self._internal or InternalBridgeHttpClient.from_runtime_config(self.runtime_config)
        if internal is None:
            raise ValueError("巡检数据处理需要配置内网HTTP桥接地址")
        today = datetime.now().strftime("%Y-%m-%d")
        deadline = time.monotonic() + max(0, float(self.cfg["wait_source_timeout_sec"]))
        if source_date != today:
            deadline = time.monotonic()
        required = [(building, family) for building in self.cfg["buildings"] for family in SOURCE_FAMILIES]
        dispatched = set()
        last_diagnostic = ""
        last_error_log = -60.0
        while True:
            cancel()
            try:
                results = internal.source_index_batch([
                    {"source_family": family, "bucket_or_date": source_date,
                     "building": b, "bucket_kind": "latest", "status": "ready", "limit": 24}
                    for b, family in required
                ], default_limit=24)
            except Exception as exc:
                if not InternalBridgeHttpClient._is_retryable_error(exc):
                    raise
                results = []
                if time.monotonic() - last_error_log >= 60:
                    log(f"[巡检数据处理] 内网暂不可用，保留原数据并等待: {exc}")
                    last_error_log = time.monotonic()
            selected = {b: {} for b in self.cfg["buildings"]}
            for i, (building, family) in enumerate(required):
                result = results[i] if i < len(results) else {}
                candidates = []
                for entry in result.get("entries", []) if result.get("ok") else []:
                    if (entry.get("building") != building or entry.get("source_family") != family
                            or entry.get("bucket_kind") != "latest"
                            or entry.get("status") != "ready" or entry.get("file_verified") is False):
                        continue
                    stamp = _text(entry.get("downloaded_at") or entry.get("updated_at"))
                    bucket = _text(entry.get("bucket_key"))
                    if not re.fullmatch(r"\d{4}-\d{2}-\d{2} (?:[01]\d|2[0-3])", bucket):
                        continue
                    if bucket[:10] != source_date or not stamp.startswith(source_date):
                        continue
                    try:
                        captured = datetime.strptime(stamp[:19], "%Y-%m-%d %H:%M:%S")
                    except ValueError:
                        continue
                    age = (datetime.now() - captured).total_seconds()
                    if source_date == today and (age < -300 or age > max(1, int(self.cfg["max_source_age_minutes"])) * 60):
                        continue
                    path = self._entry_path(entry)
                    candidates.append((stamp, str(entry.get("entry_id", "")), path, bucket))
                if candidates:
                    stamp, _, path, bucket = max(candidates, key=lambda c: c[:2])
                    if path is not None:
                        selected[building][family] = {
                            "path": path, "timestamp": stamp, "bucket_key": bucket,
                        }
            missing = [(b, family) for b, family in required if family not in selected[b]]
            if not missing:
                paths = [selected[b][family]["path"].resolve() for b, family in required]
                if len(set(paths)) != len(required):
                    raise ValueError("巡检需要D/E楼电气、暖通四份独立源文件，不能复用同一文件；未清空多维表")
                return selected
            missing_labels = ','.join(f"{b}/{SOURCE_FAMILIES[family]['label']}" for b, family in missing)
            if time.monotonic() >= deadline or datetime.now().strftime("%Y-%m-%d") != source_date:
                raise RuntimeError(f"{source_date} 巡检源文件缺失、过期或不可访问: {missing_labels}；未清空多维表")
            bucket_key = datetime.now().strftime("%Y-%m-%d %H")
            for family in SOURCE_FAMILIES:
                buildings = [b for b, f in missing if f == family and (b, f, bucket_key) not in dispatched]
                if not buildings:
                    continue
                cancel()
                try:
                    internal.refresh_latest_source_cache(source_family=family, buildings=buildings,
                                                         target_bucket_key=bucket_key)
                    dispatched.update((b, family, bucket_key) for b in buildings)
                except Exception as exc:
                    if not InternalBridgeHttpClient._is_retryable_error(exc):
                        raise
            message = f"[巡检数据处理] 等待内网源文件: date={source_date}, sources={missing_labels}"
            if message != last_diagnostic:
                log(message)
                last_diagnostic = message
            wake = time.monotonic() + min(max(1, float(self.cfg["wait_source_poll_sec"])), max(0, deadline-time.monotonic()))
            while time.monotonic() < wake:
                cancel()
                time.sleep(max(0, min(0.25, wake-time.monotonic())))

    def run(self, *, log: Callable[[str], None] | None = None, check_cancel=None,
            source_date=None, dry_run: bool = False) -> dict[str, Any]:
        log = log or (lambda _: None)
        cancel = check_cancel or (lambda: None)
        if not self.cfg["enabled"]:
            return {"status": "skipped", "reason": "disabled"}
        if set(self.cfg["buildings"]) != {"D楼", "E楼"} or len(self.cfg["buildings"]) != 2:
            raise ValueError("巡检数据处理仅支持D楼、E楼，必须同时预检")
        if len(set(self.cfg["tables"][b] for b in self.cfg["buildings"])) != 2:
            raise ValueError("D/E楼巡检目标表不能相同")
        date = self.normalize_source_date(source_date)
        sources = self._sources(date, log, cancel)
        parsed = {}
        catalog = load_catalog()
        for building in self.cfg["buildings"]:
            combined = {"rows": [], "valid_points": 0, "matched_devices": 0,
                        "missing_devices": 0, "diagnostics": [], "sources": {}}
            for family, contract in SOURCE_FAMILIES.items():
                cancel()
                entry = sources[building][family]
                p = self.parse_workbook(
                    entry["path"], building=building,
                    catalog=[d for d in catalog if d["discipline"] == contract["discipline"]],
                )
                source_time = int(datetime.strptime(entry["timestamp"][:19], "%Y-%m-%d %H:%M:%S")
                                  .replace(tzinfo=LOCAL_TZ).timestamp() * 1000)
                combined["rows"].extend({**r, "源文件时间": source_time, "源文件": entry["path"].name}
                                        for r in p["rows"])
                for key in ("valid_points", "matched_devices", "missing_devices"):
                    combined[key] += p[key]
                combined["diagnostics"].extend(f"{contract['label']}: {item}" for item in p["diagnostics"])
                combined["sources"][family] = {
                    **contract, "source_family": family, "source_file": entry["path"].name,
                    "source_time": entry["timestamp"], "bucket_kind": "latest", "bucket_key": entry["bucket_key"],
                    **{key: value for key, value in p.items() if key != "rows"},
                }
                log(f"[巡检数据处理] {building}/{contract['label']} 设备={len(p['rows'])}, 有效测点={p['valid_points']}, 缺数据设备={p['missing_devices']}, 第一个时刻={p['value_time'] or 'E列'}, 数据列={p['value_column']}")
                for item in p["diagnostics"][:10]:
                    log(f"[巡检数据处理] {building}/{contract['label']} {item}")
            parsed[building] = combined
        source_metadata = {"source_families": copy.deepcopy(DEFAULT_CONFIG["source_families"]),
                           "source_file_count": sum(len(p["sources"]) for p in parsed.values())}
        if dry_run:
            return {"status": "success", "dry_run": True, "source_date": date,
                    **source_metadata,
                    "buildings": list(parsed), "uploaded_count": 0,
                    "planned_count": sum(len(p["rows"]) for p in parsed.values()),
                    "results": {b: {k: v for k, v in p.items() if k != "rows"} for b, p in parsed.items()}}
        client = self.make_bitable_client()
        # Preflight both tables before the first destructive operation.
        schema = self.ensure_schema()
        if any(item["missing_fields"] for item in schema.values()):
            raise ValueError("巡检目标表字段未创建完整，请先执行字段初始化；未清空多维表")
        replacements = {}
        # Validate old IDs and preserve both snapshots before clearing either table.
        for building in parsed:
            cancel()
            table = self.cfg["tables"][building]
            old = client.list_records(table, page_size=500, max_records=0)
            ids = [_text(r.get("record_id")) for r in old]
            if any(not rid for rid in ids) or len(ids) != len(set(ids)):
                raise RuntimeError(f"{building} 旧记录ID不完整，未清空目标表")
            backup = BACKUP_ROOT / f"{table}.json"
            snapshot = json.loads(backup.read_text(encoding="utf-8")) if backup.exists() else {}
            # A retry must not replace the pre-failure snapshot with a half-written table.
            if snapshot.get("status") != "pending":
                snapshot = {"building": building, "status": "pending", "saved_at": datetime.now().isoformat(), "records": old}
                atomic_write_text(backup, json.dumps(snapshot, ensure_ascii=False))
            replacements[building] = (table, ids, backup, snapshot)
        result = {}
        upload_time = int(datetime.now(LOCAL_TZ).timestamp() * 1000)
        for building, parsed_building in parsed.items():
            table, ids, backup, snapshot = replacements[building]
            rows = [{**r, "上传日期": upload_time} for r in parsed_building["rows"]]
            cancel()
            try:
                if ids and client.batch_delete_records(table, ids, batch_size=200) != len(ids):
                    raise RuntimeError("清空旧记录数量与预期不符")
                responses = client.batch_create_records(table, rows, batch_size=200)
                created = [r.get("record_id") for response in responses for r in response.get("data", {}).get("records", [])]
                if len(created) != len(rows) or not all(created) or len(set(created)) != len(created):
                    raise RuntimeError("新记录创建返回数量不完整")
                identity_fields = ["设备编号", "位置", "楼栋", "专业", "源文件", "源文件时间", "上传日期", "数据完整性"]
                actual = client.list_records(table, page_size=500, max_records=0, field_names=identity_fields)
                if {r.get("record_id") for r in actual} != set(created):
                    raise RuntimeError("上传后回查数量或记录ID不一致")
                expected = {tuple(r[key] for key in identity_fields) for r in rows}
                observed = {tuple(r.get("fields", {}).get(key) for key in identity_fields) for r in actual}
                if observed != expected:
                    raise RuntimeError("上传后回查设备、位置或上传批次不一致")
                snapshot["status"] = "completed"
                atomic_write_text(backup, json.dumps(snapshot, ensure_ascii=False))
            except Exception as exc:
                raise RuntimeError(f"{building} 巡检表替换失败: {exc}；旧记录快照保存在{backup}，可重跑覆盖恢复") from exc
            result[building] = {"table_id": table, "uploaded_count": len(rows), "deleted_count": len(ids),
                                "missing_devices": parsed_building["missing_devices"],
                                "unmatched_or_invalid_rows": len(parsed_building["diagnostics"]),
                                "sources": parsed_building["sources"]}
            log(f"[巡检数据处理] {building} 上传完成: records={len(rows)}, deleted={len(ids)}")
        return {"status": "success", "source_date": date, "buildings": list(result), "results": result,
                **source_metadata,
                "uploaded_count": sum(v["uploaded_count"] for v in result.values())}
