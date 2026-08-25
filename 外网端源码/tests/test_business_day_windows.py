from __future__ import annotations

from datetime import datetime
from types import SimpleNamespace

import openpyxl

from app.modules.report_pipeline.service.calc_io_runtime import build_results_from_file_items
from app.modules.report_pipeline.service.calc_source_runtime import _business_day_data_columns
from app.modules.scheduler.service.apscheduler_orchestrator import ApschedulerSchedulerFacade


def test_monthly_columns_keep_business_day_and_next_midnight_only():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    headers = [
        "2026-08-23 23:50:00",
        "2026-08-24 00:00:00",
        "2026-08-24 12:00:00",
        "2026-08-25 00:00:00",
        "2026-08-25 00:10:00",
    ]
    for col, value in enumerate(headers, start=5):
        worksheet.cell(2, col).value = value

    columns = _business_day_data_columns(
        worksheet,
        header_row=2,
        data_start_col=5,
        data_end_col=9,
        business_date="2026-08-24",
    )

    assert columns == [6, 7, 8]


def test_file_item_business_date_reaches_monthly_calculator(tmp_path):
    source_file = tmp_path / "A楼.xlsx"
    openpyxl.Workbook().save(source_file)
    calls = []

    def calculate(file_path, building, business_date):
        calls.append((file_path, building, business_date))
        return SimpleNamespace(missing_metrics=[])

    build_results_from_file_items(
        [{"building": "A楼", "file_path": str(source_file), "upload_date": "2026-08-24"}],
        calculate_monthly_report=calculate,
        emit_log=lambda _line: None,
    )

    assert calls == [(str(source_file), "A楼", "2026-08-24")]


def test_daily_interval_offset_240_means_0400():
    facade = SimpleNamespace(
        cfg={"interval_minutes": 1440, "minute_offset": 240, "align_to_wall_clock": True},
        state={},
    )

    next_run = ApschedulerSchedulerFacade._next_interval_run_after(
        facade,
        datetime(2026, 8, 24, 3, 59, 59),
    )

    assert next_run == datetime(2026, 8, 24, 4, 0, 0)
