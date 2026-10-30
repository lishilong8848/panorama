from datetime import datetime, timedelta
import json
from pathlib import Path

import pytest
from openpyxl import Workbook

import app.modules.inspection_data.service.inspection_data_upload_service as module
from app.modules.inspection_data.service.inspection_data_upload_service import (
    InspectionDataUploadService, load_catalog, field_definitions,
)


def source(path, rows):
    wb = Workbook()
    for row in rows:
        wb.active.append([None, *row])
    wb.save(path)
    wb.close()
    return path


@pytest.fixture(autouse=True)
def forbid_live_clients(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail('Tests must use mock clients, not live bridge or Feishu access')
    monkeypatch.setattr(module.InternalBridgeHttpClient, 'from_runtime_config', forbidden)
    monkeypatch.setattr(module, 'require_feishu_auth_settings', forbidden)
    monkeypatch.setattr(module, 'FeishuBitableClient', forbidden)


def test_catalog_covers_all_forms_and_e_specific_points():
    catalog = load_catalog()
    assert len({d['workbook'] for d in catalog}) == 4
    assert sum(d['building'] == 'D楼' for d in catalog) == 458
    assert sum(d['building'] == 'E楼' for d in catalog) == 521
    assert sum(len(d['points']) for d in catalog) == 2151
    assert len({(d['building'], d['position'], d['device']) for d in catalog}) == len(catalog)
    e = [d for d in catalog if d['building'] == 'E楼']
    assert any(d['device'] == '14#柴发' for d in e)
    assert any(d['device'] == 'E-279直流屏' for d in e)
    assert not any(d['device'] == '14#柴发' for d in catalog if d['building'] == 'D楼')
    assert {p['field'] for d in e for p in d['points']} <= {f['field_name'] for f in field_definitions()}


def test_merged_rows_zero_and_missing_are_distinct(tmp_path):
    p = source(tmp_path/'d.xlsx', [
        ['电气系统西区','D-144-TR-101','L1温度(℃)',0],
        [None,None,'L2温度(℃)',35],
        [None,None,'L3温度(℃)',36],
        [None,None,'告警状态','正常'],
        [None,None,'启停状态',1],
    ])
    result = InspectionDataUploadService.parse_workbook(p, building='D楼')
    row = next(r for r in result['rows'] if r['设备编号']=='D-144-TR-101')
    assert row['L1温度(℃)'] == 0
    assert row['启停状态'] == '开启'
    assert row['数据完整性'] == '完整'
    missing = next(r for r in result['rows'] if r['设备编号']=='D-144-TR-102')
    assert 'L1温度(℃)' not in missing
    assert missing['数据完整性'] == '缺少源数据'


def test_equal_crah_names_need_room_context(tmp_path):
    p = source(tmp_path/'d.xlsx', [
        ['M1包间精密空调/北侧（240）','CRAH01','回风温度',26],
        ['M1包间精密空调/南侧（241）','CRAH01','回风温度',27],
        ['不明确位置','CRAH01','回风温度',99],
    ])
    result = InspectionDataUploadService.parse_workbook(p, building='D楼')
    assert result['matched_devices'] == 2
    assert len(result['diagnostics']) == 1
    values = [r['回风温度(℃)'] for r in result['rows'] if '回风温度(℃)' in r]
    assert sorted(values) == [26,27]


def test_crah_full_source_name_uses_printed_room(tmp_path):
    p = source(tmp_path/'d.xlsx', [['240','D-240-CRAH-01','回风温度(℃)',25]])
    result = InspectionDataUploadService.parse_workbook(p, building='D楼')
    assert result['matched_devices'] == 1
    assert next(r for r in result['rows'] if '回风温度(℃)' in r)['位置'].endswith('北侧（240）')


def test_generic_phase_temperature_is_not_guessed(tmp_path):
    p=source(tmp_path/'d.xlsx', [['西区','D-144-TR-101','温度',55]])
    with pytest.raises(ValueError, match='没有可识别'):
        InspectionDataUploadService.parse_workbook(p,building='D楼')


@pytest.mark.parametrize('value',['nan','inf','-inf','2A'])
def test_invalid_temperature_not_converted_to_zero(tmp_path,value):
    p=source(tmp_path/'d.xlsx', [['西区','D-144-TR-101','L1温度(℃)',value]])
    with pytest.raises(ValueError,match='没有可识别'):
        InspectionDataUploadService.parse_workbook(p,building='D楼')


def test_wrong_building_and_conflicting_duplicate_rejected(tmp_path):
    p=source(tmp_path/'d.xlsx', [['E楼','E-144-TR-101','L1温度(℃)',35]])
    with pytest.raises(ValueError,match='其他楼栋'):
        InspectionDataUploadService.parse_workbook(p,building='D楼')
    p=source(p,[['西区','D-144-TR-101','L1温度(℃)',35],[None,None,'L1温度(℃)',36]])
    with pytest.raises(ValueError,match='重复点位'):
        InspectionDataUploadService.parse_workbook(p,building='D楼')


def test_alarm_health_does_not_imply_running(tmp_path):
    p=source(tmp_path/'d.xlsx', [['西区','D-144-TR-101','告警状态','正常']])
    row=next(r for r in InspectionDataUploadService.parse_workbook(p,building='D楼')['rows'] if r['设备编号']=='D-144-TR-101')
    assert '启停状态' not in row


class Cloud:
    def __init__(self):
        self.rows={t: [{'record_id':'old'+t,'fields':{'设备编号':'old'}}] for t in module.DEFAULT_CONFIG['tables'].values()}
        self.events=[]
        self.fail=False
    def list_fields(self,table):
        return field_definitions()
    def list_records(self,table,**kwargs):
        assert kwargs['max_records']==0
        return self.rows[table]
    def batch_delete_records(self,table,ids,**kwargs):
        self.events.append(('delete',table))
        self.rows[table]=[r for r in self.rows[table] if r['record_id'] not in ids]
        return len(ids)
    def batch_create_records(self,table,rows,**kwargs):
        self.events.append(('create',table))
        if self.fail:
            raise RuntimeError('network timeout')
        created=[{'record_id':f'{table}-{i}','fields':r} for i,r in enumerate(rows)]
        self.rows[table].extend(created)
        return [{'data':{'records':created}}]


def inspection_files(tmp_path, invalid_source=None):
    files = {}
    catalog = load_catalog()
    for building in ['D楼', 'E楼']:
        files[building] = {}
        for index, (family, contract) in enumerate(module.SOURCE_FAMILIES.items()):
            device = next(d for d in catalog if d['building'] == building and d['discipline'] == contract['discipline'])
            point = device['points'][0]
            value = 36 if point['type'] == 2 else '原始状态'
            if invalid_source == (building, family):
                value = None
            path = source(tmp_path/f'{building}-{family}.xlsx', [
                ['位置', '设备编号', '指标', '最大值', f'2026-10-27 09:{40+index*5}:00', '2026-10-27 09:50:00'],
                [device['position'], device['device'], point['field'], 99, value, 88],
            ])
            files[building][family] = {'path': path, 'timestamp': f'2026-10-27 10:0{index*5}:00',
                                       'bucket_key': '2026-10-27 10'}
    return files


def prepared(tmp_path,monkeypatch,invalid_e=False,invalid_source=None):
    if invalid_e:
        invalid_source = ('E楼', 'electrical_inspection_family')
    files = inspection_files(tmp_path, invalid_source)
    cloud=Cloud()
    service=InspectionDataUploadService({},bitable_client=cloud)
    monkeypatch.setattr(service,'_sources',lambda *_:files)
    monkeypatch.setattr(module,'BACKUP_ROOT',tmp_path/'backups')
    backups=[]
    monkeypatch.setattr(module,'atomic_write_text',lambda path,text:backups.append((path,text)))
    return service,cloud,backups


def test_preflight_all_sources_before_clearing(tmp_path,monkeypatch):
    service,cloud,_=prepared(tmp_path,monkeypatch,invalid_e=True)
    with pytest.raises(ValueError,match='没有可识别'):
        service.run()
    assert cloud.events==[]


def test_replace_both_tables_and_retry_does_not_append(tmp_path,monkeypatch):
    service,cloud,backups=prepared(tmp_path,monkeypatch)
    for _ in range(2):
        result=service.run()
        assert result['uploaded_count']==979
        assert len(cloud.rows[module.DEFAULT_CONFIG['tables']['D楼']])==458
        assert len(cloud.rows[module.DEFAULT_CONFIG['tables']['E楼']])==521
    assert [e[0] for e in cloud.events]==['delete','create']*4
    assert len(backups)==8
    row=cloud.rows[module.DEFAULT_CONFIG['tables']['D楼']][0]['fields']
    assert isinstance(row['上传日期'],int) and row['源文件时间']>0


def test_dry_run_and_cancellation_never_delete(tmp_path,monkeypatch):
    service,cloud,backups=prepared(tmp_path,monkeypatch)
    assert service.run(dry_run=True)['dry_run'] is True
    assert cloud.events==[] and backups==[]
    def cancel():
        raise RuntimeError('cancelled')
    with pytest.raises(RuntimeError,match='cancelled'):
        service.run(check_cancel=cancel)
    assert cloud.events==[]


def test_create_failure_is_not_success_and_has_backup(tmp_path,monkeypatch):
    service,cloud,backups=prepared(tmp_path,monkeypatch)
    cloud.fail=True
    with pytest.raises(RuntimeError,match='旧记录快照'):
        service.run()
    assert len(backups)==2
    assert len(cloud.rows[module.DEFAULT_CONFIG['tables']['E楼']])==1


def test_missing_file_does_not_accept_ready_index(tmp_path):
    class Internal:
        def source_index_batch(self,queries,**kwargs):
            stamp=datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            return [{'ok':True,'entries':[{'building':q['building'],'source_family':q['source_family'],
                    'status':'ready','bucket_kind':'latest','bucket_key':stamp[:13],'downloaded_at':stamp,'file_path':str(tmp_path/'gone.xlsx')}]} for q in queries]
    cloud=Cloud()
    service=InspectionDataUploadService({'inspection_data_upload':{'wait_source_timeout_sec':0}},internal_client=Internal(),bitable_client=cloud)
    with pytest.raises(RuntimeError,match='未清空多维表'):
        service.run()
    assert cloud.events==[]


@pytest.mark.parametrize('building',['D楼','E楼'])
@pytest.mark.parametrize('discipline',['电气','暖通'])
def test_every_catalog_point_round_trips_without_missing_or_cross_device_values(tmp_path,building,discipline):
    catalog=[d for d in load_catalog() if d['building']==building and d['discipline']==discipline]
    inputs=[]
    for i,d in enumerate(catalog):
        for p in d['points']:
            inputs.append([d['position'],d['device'],p['field'],i+0.25 if p['type']==2 else '原始状态'])
        inputs.append([d['position'],d['device'],'启停状态',i%2])
    result=InspectionDataUploadService.parse_workbook(source(tmp_path/'full.xlsx',inputs),building=building,catalog=catalog)
    assert result['diagnostics']==[]
    assert result['missing_devices']==0
    assert result['matched_devices']==len(catalog)
    for i,(d,r) in enumerate(zip(catalog,result['rows'])):
        assert r['缺少源数据']==''
        assert r['启停状态']==('开启' if i%2 else '关闭')
        for p in d['points']:
            assert r[p['field']]==(i+0.25 if p['type']==2 else '原始状态')


def test_newest_missing_file_does_not_fall_back_to_older_file(tmp_path):
    p=source(tmp_path/'old.xlsx',[['西区','D-144-TR-101','L1温度(℃)',35]])
    class Internal:
        def source_index_batch(self,queries,**kwargs):
            stamp=datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            return [{'ok':True,'entries':[{'building':q['building'],'source_family':q['source_family'],
                'status':'ready','bucket_kind':'latest','bucket_key':stamp[:13],'downloaded_at':stamp,
                'entry_id':i,'file_path':str(path)} for i,path in [('1',p),('2',tmp_path/'new-but-gone.xlsx')]]} for q in queries]
    s=InspectionDataUploadService({'inspection_data_upload':{'wait_source_timeout_sec':0}},internal_client=Internal())
    with pytest.raises(RuntimeError,match='未清空多维表'):
        s.run()


def test_field_conflict_in_e_prevents_d_clear(tmp_path,monkeypatch):
    service,cloud,_=prepared(tmp_path,monkeypatch)
    original=cloud.list_fields
    def fields(table):
        definitions=original(table)
        if table==module.DEFAULT_CONFIG['tables']['E楼']:
            definitions[0]['type']=2
        return definitions
    cloud.list_fields=fields
    with pytest.raises(ValueError,match='类型不符'):
        service.run()
    assert cloud.events==[]


def test_same_room_changed_device_blank_does_not_reuse_previous_room(tmp_path):
    p=source(tmp_path/'d.xlsx', [['西区','D-144-TR-101','L1温度(℃)',35],['其他位置',None,'L2温度(℃)',99]])
    row=next(r for r in InspectionDataUploadService.parse_workbook(p,building='D楼')['rows'] if r['设备编号']=='D-144-TR-101')
    assert 'L2温度(℃)' not in row


def test_retry_preserves_pre_failure_snapshot(tmp_path,monkeypatch):
    original_write=module.atomic_write_text
    service,cloud,_=prepared(tmp_path,monkeypatch)
    monkeypatch.setattr(module,'atomic_write_text',original_write)
    cloud.fail=True
    for _ in range(2):
        with pytest.raises(RuntimeError):
            service.run()
    backup=tmp_path/'backups'/f"{module.DEFAULT_CONFIG['tables']['D楼']}.json"
    snapshot=json.loads(backup.read_text(encoding='utf-8'))
    assert snapshot['status']=='pending'
    assert snapshot['records'][0]['fields']['设备编号']=='old'
    cloud.fail=False
    service.run()
    assert json.loads(backup.read_text(encoding='utf-8'))['status']=='completed'


def test_fixed_first_timestamp_column_ignores_summary_and_later_values(tmp_path):
    wb=Workbook()
    s=wb.active
    s.append([None,'位置','设备编号','指标','最大值','2026-10-27 09:00:00','2026-10-27 09:05:00'])
    s.append([None,'西区','D-144-TR-101','L1温度(℃)',99,35,36])
    s.append([None,None,None,'L2温度(℃)',99,None,37])
    p=tmp_path/'times.xlsx';wb.save(p);wb.close()
    parsed=InspectionDataUploadService.parse_workbook(p,building='D楼')
    row=next(r for r in parsed['rows'] if r['设备编号']=='D-144-TR-101')
    assert row['L1温度(℃)']==35
    assert 'L2温度(℃)' not in row
    assert parsed['value_column']==6 and parsed['value_time']=='2026-10-27 09:00:00'


def test_multi_column_without_time_headers_fails_before_upload(tmp_path):
    wb=Workbook();wb.active.append([None,'西区','D-144-TR-101','L1温度(℃)',35,99]);p=tmp_path/'unknown.xlsx';wb.save(p);wb.close()
    with pytest.raises(ValueError,match='时间表头'):
        InspectionDataUploadService.parse_workbook(p,building='D楼')


SOURCE_PAIRS = [(b, family) for b in ['D楼', 'E楼'] for family in module.SOURCE_FAMILIES]


class IndexedSources:
    def __init__(self, files):
        self.queries = []
        self.refreshes = []
        stamp = module.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        self.entries = {
            (building, family): {
                'building': building, 'source_family': family, 'status': 'ready',
                'bucket_kind': 'latest', 'bucket_key': stamp[:13], 'downloaded_at': stamp,
                'file_verified': True, 'file_path': str(entry['path']),
            }
            for building, sources in files.items() for family, entry in sources.items()
        }

    def source_index_batch(self, queries, **kwargs):
        self.queries.append(queries)
        return [{'ok': True, 'entries': [self.entries[(q['building'], q['source_family'])]]
                 if (q['building'], q['source_family']) in self.entries else []} for q in queries]

    def refresh_latest_source_cache(self, **kwargs):
        self.refreshes.append(kwargs)


@pytest.mark.parametrize('pair', SOURCE_PAIRS)
def test_each_of_four_sources_must_parse_before_any_clear(tmp_path, monkeypatch, pair):
    service, cloud, backups = prepared(tmp_path, monkeypatch, invalid_source=pair)
    with pytest.raises(ValueError, match='没有可识别'):
        service.run()
    assert cloud.events == [] and backups == []


@pytest.mark.parametrize('pair', SOURCE_PAIRS)
def test_each_missing_source_is_named_and_blocks_clear(tmp_path, pair):
    internal = IndexedSources(inspection_files(tmp_path))
    del internal.entries[pair]
    cloud = Cloud()
    service = InspectionDataUploadService({'inspection_data_upload': {'wait_source_timeout_sec': 0}},
                                         internal_client=internal, bitable_client=cloud)
    with pytest.raises(RuntimeError, match=f'{pair[0]}/{module.SOURCE_FAMILIES[pair[1]]["label"]}'):
        service.run()
    assert cloud.events == [] and internal.refreshes == []
    queries = internal.queries[0]
    assert {(q['building'], q['source_family']) for q in queries} == set(SOURCE_PAIRS)
    assert all(q['bucket_kind'] == 'latest' for q in queries)


@pytest.mark.parametrize('patch', [
    {'source_family': 'inspection_data_family'},
    {'source_family': 'electrical_inspection_family'},
    {'building': 'D楼'},
    {'bucket_kind': 'daily'},
    {'bucket_key': '2026-10-30'},
    {'bucket_key': '2026-10-30 99'},
    {'file_verified': False},
    {'status': 'failed'},
    {'downloaded_at': 'invalid'},
])
def test_source_index_requires_explicit_matching_contract(tmp_path, patch):
    internal = IndexedSources(inspection_files(tmp_path))
    internal.entries[('E楼', 'hvac_inspection_family')].update(patch)
    cloud = Cloud()
    service = InspectionDataUploadService({'inspection_data_upload': {'wait_source_timeout_sec': 0}},
                                         internal_client=internal, bitable_client=cloud)
    with pytest.raises(RuntimeError, match='E楼/暖通巡检源文件'):
        service.run()
    assert cloud.events == []


def test_legacy_combined_files_are_never_a_fallback(tmp_path):
    internal = IndexedSources(inspection_files(tmp_path))
    for entry in internal.entries.values():
        entry['source_family'] = 'inspection_data_family'
    service = InspectionDataUploadService({'inspection_data_upload': {'wait_source_timeout_sec': 0}},
                                         internal_client=internal, bitable_client=Cloud())
    with pytest.raises(RuntimeError, match='未清空多维表'):
        service.run(dry_run=True)


def test_four_index_entries_cannot_reuse_a_combined_file(tmp_path):
    internal = IndexedSources(inspection_files(tmp_path))
    for building in ['D楼', 'E楼']:
        internal.entries[(building, 'hvac_inspection_family')]['file_path'] = internal.entries[(building, 'electrical_inspection_family')]['file_path']
    service = InspectionDataUploadService({}, internal_client=internal, bitable_client=Cloud())
    with pytest.raises(ValueError, match='四份独立源文件'):
        service.run(dry_run=True)


@pytest.mark.parametrize('pair', SOURCE_PAIRS)
def test_valid_points_from_other_discipline_cannot_satisfy_preflight(tmp_path, monkeypatch, pair):
    service, cloud, backups = prepared(tmp_path, monkeypatch)
    building, family = pair
    other = next(d for d in load_catalog() if d['building'] == building and d['discipline'] != module.SOURCE_FAMILIES[family]['discipline'])
    point = other['points'][0]
    source(tmp_path/f'{building}-{family}.xlsx', [[other['position'], other['device'], point['field'], 55]])
    with pytest.raises(ValueError, match='没有可识别'):
        service.run()
    assert cloud.events == [] and backups == []


@pytest.mark.parametrize('dry_run', [False, True])
def test_split_result_reports_four_files_and_both_timestamps(tmp_path, monkeypatch, dry_run):
    service, cloud, _ = prepared(tmp_path, monkeypatch)
    result = service.run(dry_run=dry_run)
    assert result['source_file_count'] == 4
    assert result['source_families'] == {'电气': 'electrical_inspection_family', '暖通': 'hvac_inspection_family'}
    assert 'source_family' not in result
    assert result['uploaded_count'] == (0 if dry_run else 979)
    for building, details in result['results'].items():
        assert set(details['sources']) == set(module.SOURCE_FAMILIES)
        assert 'source_time' not in details and 'value_time' not in details
        for index, (family, contract) in enumerate(module.SOURCE_FAMILIES.items()):
            metadata = details['sources'][family]
            assert metadata['report_name'] == contract['report_name']
            assert metadata['label'] == contract['label']
            assert metadata['source_family'] == family
            assert metadata['source_file'] == f'{building}-{family}.xlsx'
            assert metadata['source_time'] == f'2026-10-27 10:0{index*5}:00'
            assert metadata['value_column'] == 6
            assert metadata['value_time'] == f'2026-10-27 09:{40+index*5}:00'
            assert metadata['matched_devices'] == 1
            if not dry_run:
                rows = [r['fields'] for r in cloud.rows[module.DEFAULT_CONFIG['tables'][building]]]
                rows = [r for r in rows if r['专业'] == contract['discipline']]
                assert rows and all(r['源文件'] == metadata['source_file'] for r in rows)
                expected_time = int(datetime.strptime(metadata['source_time'], '%Y-%m-%d %H:%M:%S').replace(tzinfo=module.LOCAL_TZ).timestamp()*1000)
                assert all(r['源文件时间'] == expected_time for r in rows)
    if dry_run:
        assert cloud.events == []


def test_fixed_family_mapping_ignores_legacy_and_crossmapped_config(tmp_path):
    internal = IndexedSources(inspection_files(tmp_path))
    service = InspectionDataUploadService({'inspection_data_upload': {
        'source_family': 'inspection_data_family',
        'source_families': {'电气': 'hvac_inspection_family', '暖通': 'electrical_inspection_family'},
    }}, internal_client=internal, bitable_client=Cloud())
    result = service.run(dry_run=True)
    assert 'source_family' not in service.cfg
    assert result['source_families'] == module.DEFAULT_CONFIG['source_families']
    assert {(q['building'], q['source_family']) for q in internal.queries[0]} == set(SOURCE_PAIRS)


def test_other_discipline_readings_do_not_fill_missing_devices(tmp_path, monkeypatch):
    service, cloud, _ = prepared(tmp_path, monkeypatch)
    catalog = load_catalog()
    for building, family in SOURCE_PAIRS:
        discipline = module.SOURCE_FAMILIES[family]['discipline']
        own = next(d for d in catalog if d['building'] == building and d['discipline'] == discipline)
        other = [d for d in catalog if d['building'] == building and d['discipline'] != discipline][1]
        source(tmp_path/f'{building}-{family}.xlsx', [
            [own['position'], own['device'], own['points'][0]['field'], 36],
            [other['position'], other['device'], other['points'][0]['field'], 99],
        ])
    service.run()
    for building in ['D楼', 'E楼']:
        rows = [r['fields'] for r in cloud.rows[module.DEFAULT_CONFIG['tables'][building]]]
        for discipline in ['电气', '暖通']:
            device = [d for d in catalog if d['building'] == building and d['discipline'] == discipline][1]
            row = next(r for r in rows if r['专业'] == discipline and r['设备编号'] == device['device'] and r['位置'] == device['position'])
            assert all(p['field'] not in row for p in device['points'])
            assert '启停状态' not in row
            assert row['数据完整性'] == '缺少源数据'


def test_e_old_record_preflight_failure_prevents_d_clear(tmp_path, monkeypatch):
    service, cloud, _ = prepared(tmp_path, monkeypatch)
    cloud.rows[module.DEFAULT_CONFIG['tables']['E楼']] = [{'fields': {'设备编号': 'old'}}]
    with pytest.raises(RuntimeError, match='旧记录ID不完整'):
        service.run()
    assert cloud.events == []


def test_e_backup_failure_prevents_d_clear(tmp_path, monkeypatch):
    service, cloud, _ = prepared(tmp_path, monkeypatch)
    original = module.atomic_write_text
    def write(path, content):
        if path.stem == module.DEFAULT_CONFIG['tables']['E楼']:
            raise OSError('disk full')
        original(path, content)
    monkeypatch.setattr(module, 'atomic_write_text', write)
    with pytest.raises(OSError, match='disk full'):
        service.run()
    assert cloud.events == []


@pytest.fixture
def polling_clock(monkeypatch):
    class Clock:
        elapsed = 0.0
        def monotonic(self):
            return self.elapsed
        def sleep(self, seconds):
            self.elapsed += seconds
    clock = Clock()
    class WallClock(datetime):
        @classmethod
        def now(cls, tz=None):
            value = cls(2026, 10, 30, 10, 0, 0) + timedelta(seconds=clock.elapsed)
            return value.replace(tzinfo=tz) if tz else value
    monkeypatch.setattr(module, 'time', clock)
    monkeypatch.setattr(module, 'datetime', WallClock)
    return clock


def test_refresh_retries_only_missing_family_and_building(tmp_path, polling_clock):
    internal = IndexedSources(inspection_files(tmp_path))
    missing = [('D楼', 'electrical_inspection_family'), ('E楼', 'hvac_inspection_family')]
    saved = {pair: internal.entries.pop(pair) for pair in missing}
    def refresh(**kwargs):
        internal.refreshes.append(kwargs)
        family = kwargs['source_family']
        if family == 'hvac_inspection_family' and sum(r['source_family'] == family for r in internal.refreshes) == 1:
            raise RuntimeError('connection refused')
        for building in kwargs['buildings']:
            internal.entries[(building, family)] = saved[(building, family)]
    internal.refresh_latest_source_cache = refresh
    service = InspectionDataUploadService({'inspection_data_upload': {'wait_source_timeout_sec': 5, 'wait_source_poll_sec': 1}},
                                         internal_client=internal, bitable_client=Cloud())
    result = service.run(dry_run=True)
    assert result['source_file_count'] == 4
    assert [(r['source_family'], r['buildings']) for r in internal.refreshes] == [
        ('electrical_inspection_family', ['D楼']), ('hvac_inspection_family', ['E楼']), ('hvac_inspection_family', ['E楼']),
    ]
    assert all(r['target_bucket_key'] == '2026-10-30 10' for r in internal.refreshes)


def test_accepted_refresh_is_not_repeated_while_waiting(tmp_path, polling_clock):
    internal = IndexedSources(inspection_files(tmp_path))
    internal.entries.pop(('E楼', 'hvac_inspection_family'))
    service = InspectionDataUploadService({'inspection_data_upload': {'wait_source_timeout_sec': 3, 'wait_source_poll_sec': 1}},
                                         internal_client=internal, bitable_client=Cloud())
    with pytest.raises(RuntimeError, match='E楼/暖通巡检源文件'):
        service.run()
    assert internal.refreshes == [{'source_family': 'hvac_inspection_family', 'buildings': ['E楼'], 'target_bucket_key': '2026-10-30 10'}]


def test_historical_missing_source_does_not_request_current_hour(tmp_path, polling_clock):
    internal = IndexedSources(inspection_files(tmp_path))
    for entry in internal.entries.values():
        entry.update(bucket_key='2026-10-29 10', downloaded_at='2026-10-29 10:00:00')
    internal.entries.pop(('E楼', 'hvac_inspection_family'))
    service = InspectionDataUploadService({}, internal_client=internal, bitable_client=Cloud())
    with pytest.raises(RuntimeError, match='E楼/暖通巡检源文件'):
        service.run(source_date='2026-10-29')
    assert internal.refreshes == []


def test_stale_source_refreshes_only_its_pair(tmp_path, polling_clock):
    internal = IndexedSources(inspection_files(tmp_path))
    entry = internal.entries[('E楼', 'hvac_inspection_family')]
    entry.update(bucket_key='2026-10-30 07', downloaded_at='2026-10-30 07:00:00')
    def refresh(**kwargs):
        internal.refreshes.append(kwargs)
        entry.update(bucket_key='2026-10-30 10', downloaded_at='2026-10-30 10:00:00')
    internal.refresh_latest_source_cache = refresh
    service = InspectionDataUploadService({}, internal_client=internal, bitable_client=Cloud())
    assert service.run(dry_run=True)['source_file_count'] == 4
    assert len(internal.refreshes) == 1
    assert internal.refreshes[0]['source_family'] == 'hvac_inspection_family'
    assert internal.refreshes[0]['buildings'] == ['E楼']


def test_wait_crossing_midnight_never_refreshes_wrong_date(tmp_path, polling_clock):
    polling_clock.elapsed = 14*60*60 - 1
    internal = IndexedSources(inspection_files(tmp_path))
    internal.entries.pop(('E楼', 'hvac_inspection_family'))
    service = InspectionDataUploadService({'inspection_data_upload': {'wait_source_timeout_sec': 3, 'wait_source_poll_sec': 1}},
                                         internal_client=internal, bitable_client=Cloud())
    with pytest.raises(RuntimeError, match='E楼/暖通巡检源文件'):
        service.run()
    assert internal.refreshes == [{'source_family': 'hvac_inspection_family', 'buildings': ['E楼'], 'target_bucket_key': '2026-10-30 23'}]


@pytest.mark.parametrize('building', ['D楼', 'E楼'])
def test_explicit_own_building_label_is_accepted_and_other_building_rejected(tmp_path, building):
    path = source(tmp_path/'labeled.xlsx', [[f'{building} 西区', f'{building[0]}-144-TR-101', 'L1温度(℃)', 35]])
    result = InspectionDataUploadService.parse_workbook(path, building=building)
    assert result['matched_devices'] == 1
    other = 'D楼' if building == 'E楼' else 'E楼'
    with pytest.raises(ValueError, match='其他楼栋'):
        InspectionDataUploadService.parse_workbook(path, building=other)
