import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from handover_log_module.repository.excel_reader import load_workbook_quietly
from handover_log_module.repository.review_building_document_store import ReviewBuildingDocumentStore
from handover_log_module.service.capacity_room_inputs_service import CapacityRoomInputsService
from handover_log_module.service.handover_capacity_report_service import HandoverCapacityReportService
from handover_log_module.service.review_document_state_service import ReviewDocumentStateError, ReviewDocumentStateService


def _document():
    cells = dict(zip(("AA69", "AA79", "AA89", "AA103", "AA117", "AA127"), ("168", "170", "140", "152", "170", "170")))
    cells.update(J159="7", J167="8", J175="9")
    return {
        "fixed_blocks": [{"fields": [
            {"cell": cell, "value": value} for cell, value in {
                "B13": "988", "D13": "970", "F13": "0", "H13": "0", "H6": "24",
                "F8": "西区30/东区30", "B6": "1.2", "D6": "1200", "F6": "1000",
            }.items()
        ]}],
        "capacity_room_inputs": CapacityRoomInputsService.payload_from_cells(cells, building="A楼"),
    }


def _fixed(document, cell):
    return next(field for block in document["fixed_blocks"] for field in block["fields"] if field["cell"] == cell)


def test_a_total_is_derived_without_accumulating_or_changing_other_buildings():
    document = _document()
    updated, changed = CapacityRoomInputsService.apply_powered_total_to_document(document, building="A楼")
    assert changed and _fixed(updated, "D13")["value"] == "994"
    assert _fixed(updated, "D13")["read_only"] is True
    assert _fixed(document, "D13")["value"] == "970"
    assert _fixed(updated, "B13")["value"] == "988"
    again, changed = CapacityRoomInputsService.apply_powered_total_to_document(updated, building="A楼")
    assert not changed and _fixed(again, "D13")["value"] == "994"
    for building in ("B楼", "C楼", "D楼", "E楼"):
        unchanged, changed = CapacityRoomInputsService.apply_powered_total_to_document(document, building=building)
        assert not changed and unchanged == document


def test_zero_potential_rooms_and_incomplete_legacy_counts():
    cells = CapacityRoomInputsService.extract_cells_from_document(_document(), building="A楼")
    for cell in ("J159", "J167", "J175"):
        cells[cell] = "0"
    assert CapacityRoomInputsService.powered_total_from_cells(cells, building="A楼") == "970"
    cells["J159"] = "3.0"
    assert CapacityRoomInputsService.powered_total_from_cells(cells, building="A楼") == "973"
    for invalid in ("", "bad", "-1", "1.5", "NaN"):
        incomplete = {**cells, "AA69": invalid}
        assert CapacityRoomInputsService.powered_total_from_cells(incomplete, building="A楼") is None


@pytest.mark.parametrize("malformed", ["bad_count", None, {}, {"rows": []}, {"rows": "invalid"}])
def test_invalid_room_count_does_not_replace_saved_total(tmp_path, malformed):
    config = {"_global_paths": {"runtime_state_root": str(tmp_path)}}
    session = {"building": "A楼", "session_id": "A楼|2026-10-14|day", "revision": 1}
    store = ReviewBuildingDocumentStore(config=config, building="A楼")
    state = store.upsert_imported_document(session=session, document=_document(), imported_from_excel=True)
    document = _document()
    if malformed == "bad_count":
        document["capacity_room_inputs"]["rows"][0]["powered_cabinets"] = "bad"
    else:
        document["capacity_room_inputs"] = malformed
    with pytest.raises(ReviewDocumentStateError, match="非负整数"):
        ReviewDocumentStateService(config).save_document(
            session=session, document=document, base_revision=state["revision"],
            dirty_regions={"capacity_room_inputs": True}, ensure_ready=False,
        )
    assert store.get_document(session["session_id"])["document"] == state["document"]


def test_fresh_generation_uses_saved_rooms_after_restart_without_double_counting(tmp_path):
    config = {
        "_global_paths": {"runtime_state_root": str(tmp_path / "runtime")},
        "template": {"sheet_name": "Sheet"},
        "review_ui": {"cabinet_power_defaults_by_building": {"A楼": {"cells": {
            "B13": "988", "D13": "970", "F13": "0", "H13": "0",
        }}}},
    }
    ReviewDocumentStateService(config).persist_defaults_from_document(
        building="A楼", document=_document(), dirty_regions={"capacity_room_inputs": True},
    )
    output = tmp_path / "handover.xlsx"
    workbook = Workbook()
    workbook.active["D13"] = "970"
    workbook.save(output)
    workbook.close()
    for _ in range(2):
        ReviewDocumentStateService(config).apply_cabinet_defaults_to_output(
            building="A楼", output_file=output, emit_log=lambda _text: None,
        )
        checked = load_workbook(output)
        try:
            assert checked.active["D13"].value == "994"
            assert checked.active["B13"].value == "988"
        finally:
            checked.close()


def test_review_save_and_download_keep_both_workbooks_at_same_total(tmp_path, monkeypatch):
    config = {"_global_paths": {"runtime_state_root": str(tmp_path / "runtime")}, "template": {"sheet_name": "Sheet"}}
    handover = tmp_path / "handover.xlsx"
    workbook = Workbook()
    workbook.active["D13"] = "970"
    workbook.save(handover)
    workbook.close()
    session = {
        "building": "A楼", "session_id": "A楼|2026-10-14|day", "revision": 1,
        "duty_date": "2026-10-14", "duty_shift": "day", "output_file": str(handover),
    }
    state_service = ReviewDocumentStateService(config, emit_log=lambda _text: None)
    store = ReviewBuildingDocumentStore(config=config, building="A楼")
    state = store.upsert_imported_document(session=session, document=_document(), imported_from_excel=True)
    state, _ = state_service.save_document(
        session=session, document=_document(), base_revision=state["revision"],
        dirty_regions={"capacity_room_inputs": True}, ensure_ready=False,
    )
    assert _fixed(state["document"], "D13")["value"] == "994"
    assert state["dirty_regions"]["fixed_blocks"] is True
    state_service.force_sync_session(building="A楼", session_id=session["session_id"], reconcile_sync_job=False)
    checked = load_workbook(handover)
    try:
        assert checked.active["D13"].value == "994"
    finally:
        checked.close()

    capacity = tmp_path / "capacity.xlsx"
    workbook = load_workbook_quietly(Path(__file__).resolve().parents[1] / "A楼交接班容量报表空模板.xlsx")
    workbook.save(capacity)
    workbook.close()
    service = HandoverCapacityReportService(config)
    monkeypatch.setattr(service, "_query_capacity_water_summary", lambda **kwargs: {})
    monkeypatch.setattr(service, "query_total_electricity_summary", lambda **kwargs: {})
    monkeypatch.setattr(service, "_fetch_weather_payload_for_duty_date", lambda **kwargs: {})
    monkeypatch.setattr(service, "_a_building_outdoor_cells_for_capacity", lambda **kwargs: {})
    monkeypatch.setattr(service, "_build_load_rate_values_for_sheet", lambda *args, **kwargs: ({}, "", {}))
    monkeypatch.setattr(service, "build_capacity_overlay_signature", lambda **kwargs: {"signature": "offline"})
    for _ in range(2):
        tracked = service.extract_tracked_cells_from_review_document(state["document"])
        result = service.sync_overlay_for_existing_report_from_cells(
            building="A楼", duty_date="2026-10-14", duty_shift="day", handover_cells=tracked,
            capacity_output_file=str(capacity), shared_110kv={}, emit_log=lambda _text: None,
        )
        assert result["status"] == "ready", result
        checked = load_workbook(capacity)
        try:
            sheet = checked["本班组"]
            assert sheet["AC58"].value == "994"
            assert [sheet[cell].value for cell in ("J159", "J167", "J175")] == ["7", "8", "9"]
            assert sheet["AD58"].value == "=AC58/AB58"
        finally:
            checked.close()

    document = copy.deepcopy(state["document"])
    document["capacity_room_inputs"]["rows"][6]["powered_cabinets"] = "5"
    state, _ = state_service.save_document(
        session=session, document=document, base_revision=state["revision"],
        dirty_regions={"capacity_room_inputs": True}, ensure_ready=False,
    )
    assert _fixed(state["document"], "D13")["value"] == "992"
