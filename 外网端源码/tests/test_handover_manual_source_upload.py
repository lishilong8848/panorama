import asyncio
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

from app.modules.handover_review.api import routes
from app.worker import task_handlers
from handover_log_module.service.handover_orchestrator import HandoverOrchestrator, HandoverQueryContext
from handover_log_module.service.review_session_service import ReviewSessionService


def workbook_upload(name, value):
    stream = BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = value
    workbook.save(stream)
    workbook.close()
    stream.seek(0)
    return UploadFile(filename=name, file=stream)


def setup_upload(monkeypatch, tmp_path, *, active=False):
    observed = {}
    service = SimpleNamespace(
        get_session_for_building_duty_fast=lambda *_args: None,
        build_session_id=lambda b, d, s: f"{b}|{d}|{s}",
        build_batch_key=lambda d, s: f"{d}|{s}",
    )
    container = SimpleNamespace(
        job_service=SimpleNamespace(find_active_job_by_dedupe_key=lambda _key: {"job_id": "active"} if active else None),
        add_system_log=Mock(),
    )
    request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(container=container)))
    monkeypatch.setattr(routes, "get_app_dir", lambda: str(tmp_path))
    monkeypatch.setattr(routes, "_build_review_session_service", lambda _container: service)
    monkeypatch.setattr(routes, "_resolve_building_or_404", lambda *_args: "A楼")
    monkeypatch.setattr(routes, "_handover_cfg", lambda _container: {})
    monkeypatch.setattr(routes, "_resolve_regenerate_source_files", Mock(side_effect=AssertionError("must not request shared sources")))

    def start(_container, **kwargs):
        observed.update(kwargs)
        return SimpleNamespace(job_id="manual-job")

    monkeypatch.setattr(routes, "_start_handover_background_job", start)
    monkeypatch.setattr(routes, "_accepted_job_response", lambda job: {"job_id": job.job_id})
    return request, observed


def call_upload(request, handover, capacity):
    return asyncio.run(routes.handover_review_regenerate_from_files(
        "a", request, handover, capacity,
        session_id="", duty_date="2026-10-06", duty_shift="day", client_id="review-test",
    ))


def test_two_uploaded_workbooks_are_the_only_worker_sources(monkeypatch, tmp_path):
    request, observed = setup_upload(monkeypatch, tmp_path)
    handover = workbook_upload("交接班日志（李世龙）.xlsx", "log-source")
    capacity = workbook_upload("交接班容量报表.xlsx", "capacity-source")
    assert call_upload(request, handover, capacity) == {"job_id": "manual-job"}
    payload = observed["worker_payload"]
    for key, expected in (("data_file", "log-source"), ("capacity_source_file", "capacity-source")):
        path = Path(payload[key])
        assert path.is_relative_to(tmp_path)
        workbook = load_workbook(path, read_only=True)
        assert workbook.active["A1"].value == expected
        workbook.close()
    assert payload["duty_date"] == "2026-10-06"
    assert payload["duty_shift"] == "day"
    assert observed["dedupe_key"] == "handover_review_regenerate:A楼|2026-10-06|day"
    assert handover.file.closed and capacity.file.closed


@pytest.mark.parametrize("name,content", [("capacity.txt", b"text"), ("capacity.xlsx", b""), ("capacity.xlsx", b"broken")])
def test_invalid_second_file_cleans_first_and_does_not_submit(monkeypatch, tmp_path, name, content):
    request, observed = setup_upload(monkeypatch, tmp_path)
    handover = workbook_upload("log.xlsx", "valid")
    capacity = UploadFile(filename=name, file=BytesIO(content))
    with pytest.raises(HTTPException) as error:
        call_upload(request, handover, capacity)
    assert error.value.status_code == 400
    assert not observed
    assert not list(tmp_path.rglob("*.xlsx"))
    assert handover.file.closed and capacity.file.closed


def test_duplicate_generation_rejects_new_upload_without_overwriting(monkeypatch, tmp_path):
    request, observed = setup_upload(monkeypatch, tmp_path, active=True)
    with pytest.raises(HTTPException) as error:
        call_upload(request, workbook_upload("log.xlsx", "log"), workbook_upload("capacity.xlsx", "capacity"))
    assert error.value.status_code == 409
    assert "已有生成任务" in str(error.value.detail)
    assert not observed
    assert not list(tmp_path.rglob("*.xlsx"))


def test_upload_limit_has_source_specific_error():
    upload = UploadFile(filename="log.xlsx", file=BytesIO(b"12345"))
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes._read_upload_file_limited(upload, max_bytes=4, file_label="交接班日志源文件"))
    assert "交接班日志源文件" in error.value.detail
    asyncio.run(upload.close())


@pytest.mark.parametrize("method", ["run_from_existing_files", "run_from_download"])
def test_queued_automatic_job_skips_manually_generated_building_before_io(method):
    orchestrator = HandoverOrchestrator.__new__(HandoverOrchestrator)
    orchestrator.config = {"sites": [{"building": "A楼", "enabled": True}]}
    orchestrator._review_session_service = SimpleNamespace(is_manual_regenerated_for_duty=lambda **_kwargs: True)
    kwargs = {"building_files": [("A楼", "never-read.xlsx")]} if method == "run_from_existing_files" else {"buildings": ["A楼"]}
    result = getattr(orchestrator, method)(
        **kwargs, duty_date="2026-10-06", duty_shift="day", skip_manual_generated=True, emit_log=lambda _: None,
    )
    assert result["status"] == "skipped"
    assert result["reason"] == "manual_regenerated"
    assert result["skipped_buildings"] == ["A楼"]


def test_manual_marker_is_limited_to_same_building_date_and_shift():
    orchestrator = HandoverOrchestrator.__new__(HandoverOrchestrator)
    orchestrator._review_session_service = SimpleNamespace(
        is_manual_regenerated_for_duty=lambda **kw: (kw["building"], kw["duty_date"], kw["duty_shift"]) == ("A楼", "2026-10-06", "day"),
    )
    assert orchestrator._skip_manual_generated_buildings(
        ["A楼", "B楼"], duty_date="2026-10-06", duty_shift="day", emit_log=lambda _: None,
    ) == ["B楼"]
    assert orchestrator._skip_manual_generated_buildings(
        ["A楼"], duty_date="2026-10-06", duty_shift="night", emit_log=lambda _: None,
    ) == ["A楼"]
    assert orchestrator._skip_manual_generated_buildings(
        ["A楼"], duty_date="2026-10-07", duty_shift="day", emit_log=lambda _: None,
    ) == ["A楼"]


@pytest.mark.parametrize("method", ["run_from_existing_files", "run_from_download"])
@pytest.mark.parametrize("context", [{}, {"duty_date": "  ", "duty_shift": "  "}])
def test_legacy_automatic_payload_infers_missing_duty_before_manual_check(method, context):
    orchestrator = HandoverOrchestrator.__new__(HandoverOrchestrator)
    orchestrator.config = {"sites": [{"building": "A楼", "enabled": True}]}
    checker = Mock(return_value=True)
    orchestrator._review_session_service = SimpleNamespace(is_manual_regenerated_for_duty=checker)
    orchestrator._infer_duty_by_now = lambda: ("2026-10-06", "day")
    kwargs = {"building_files": [("A楼", "not-read.xlsx")]} if method == "run_from_existing_files" else {"buildings": ["A楼"]}
    result = getattr(orchestrator, method)(**kwargs, **context, skip_manual_generated=True, emit_log=lambda _: None)
    assert result["status"] == "skipped"
    checker.assert_called_once_with(building="A楼", duty_date="2026-10-06", duty_shift="day")


def test_partial_manual_completion_does_not_block_other_buildings():
    orchestrator = HandoverOrchestrator.__new__(HandoverOrchestrator)
    orchestrator.config = {"sites": [{"building": b, "enabled": True} for b in ("A楼", "B楼")]}
    orchestrator._review_session_service = SimpleNamespace(
        is_manual_regenerated_for_duty=lambda **kw: kw["building"] == "A楼",
        build_batch_key=lambda d, s: f"{d}|{s}",
    )
    orchestrator._build_query_context = lambda **kw: HandoverQueryContext(duty_date="2026-10-06", duty_shift="day", target_buildings=kw["buildings"])
    orchestrator._deployment_role_mode = lambda: "internal"
    orchestrator._build_alarm_duty_window = lambda **_kw: SimpleNamespace(start_time="start", end_time="end")
    orchestrator._build_fixed_values_with_alarm = lambda **_kw: ({}, None, {})
    orchestrator._resolve_shared_outdoor_temperature_cells = lambda *_args, **_kw: {}
    orchestrator._send_station_110_review_link_with_handover = Mock()
    orchestrator._trigger_station_h_sync_after_generation = Mock()
    orchestrator.run_from_existing_file = Mock(return_value={"results": [{"building": "B楼", "success": True}]})
    result = orchestrator.run_from_existing_files(
        building_files=[("A楼", "a.xlsx"), ("B楼", "b.xlsx")], duty_date="2026-10-06", duty_shift="day",
        skip_manual_generated=True, emit_log=lambda _: None,
    )
    assert result["selected_buildings"] == ["B楼"]
    assert result["skipped_buildings"] == ["A楼"]
    assert result["success_count"] == 1
    assert orchestrator.run_from_existing_file.call_args.kwargs["building"] == "B楼"


def test_capacity_generation_failure_does_not_mark_manual_completed(monkeypatch, tmp_path):
    source = tmp_path / "source.xlsx"
    source.touch()
    review_service = Mock()
    monkeypatch.setattr(task_handlers, "load_handover_config", lambda _cfg: {})
    monkeypatch.setattr(task_handlers, "ReviewSessionService", lambda _cfg: review_service)
    monkeypatch.setattr(task_handlers, "ReviewDocumentStateService", lambda *_args, **_kw: Mock())
    monkeypatch.setattr(task_handlers, "HandoverXlsxWriteQueueService", lambda *_args, **_kw: Mock())
    monkeypatch.setattr(task_handlers, "OrchestratorService", lambda _cfg: SimpleNamespace(
        run_handover_from_files=lambda **_kw: {"results": [{"building": "A楼", "success": True, "capacity_status": "failed", "capacity_error": "capacity invalid"}]},
    ))
    with pytest.raises(RuntimeError, match="capacity invalid"):
        task_handlers.handle_handover_review_regenerate({}, {
            "building": "A楼", "session_id": "A楼|2026-10-06|day", "duty_date": "2026-10-06", "duty_shift": "day",
            "data_file": str(source), "capacity_source_file": str(source),
        }, lambda _: None)
    review_service.mark_manual_regenerated.assert_not_called()


def test_manual_generation_marker_survives_service_restart(monkeypatch, tmp_path):
    # Production cleanup deliberately ignores pytest output directories.
    monkeypatch.setattr(ReviewSessionService, "_is_legacy_test_output_file", lambda *_args: False)
    config = {"_global_paths": {"runtime_state_root": str(tmp_path / "runtime")}}
    first = ReviewSessionService(config)
    first.configure_generated_file_index(None)
    first.register_generated_output(
        building="A楼", duty_date="2026-10-06", duty_shift="day",
        data_file=str(tmp_path / "source.xlsx"), output_file=str(tmp_path / "output.xlsx"), source_mode="from_file",
    )
    first.mark_manual_regenerated(building="A楼", duty_date="2026-10-06", duty_shift="day", client_id="review-test")
    restarted = ReviewSessionService(config)
    restarted.configure_generated_file_index(None)
    assert restarted.is_manual_regenerated_for_duty(building="A楼", duty_date="2026-10-06", duty_shift="day")
    assert not restarted.is_manual_regenerated_for_duty(building="A楼", duty_date="2026-10-06", duty_shift="night")


def test_cancelled_during_capacity_overlay_does_not_mark_completed(monkeypatch, tmp_path):
    source = tmp_path / "source.xlsx"
    source.touch()
    cancelled = [False]

    def check_cancelled():
        if cancelled[0]:
            raise RuntimeError("cancelled during overlay")

    def barrier(**kwargs):
        if kwargs["reason"] == "review_regenerate_capacity_overlay":
            cancelled[0] = True
        return {"status": "success"}

    review_service = Mock()
    review_service.get_session_by_id.return_value = {"session_id": "manual-session"}
    queue = Mock()
    queue.wait_for_barrier.side_effect = barrier
    monkeypatch.setattr(task_handlers, "load_handover_config", lambda _cfg: {})
    monkeypatch.setattr(task_handlers, "ReviewSessionService", lambda _cfg: review_service)
    monkeypatch.setattr(task_handlers, "ReviewDocumentStateService", lambda *_args, **_kw: Mock())
    monkeypatch.setattr(task_handlers, "ReviewBuildingDocumentStore", lambda **_kw: Mock())
    monkeypatch.setattr(task_handlers, "HandoverXlsxWriteQueueService", lambda *_args, **_kw: queue)
    monkeypatch.setattr(task_handlers, "OrchestratorService", lambda _cfg: SimpleNamespace(
        run_handover_from_files=lambda **_kw: {"results": [{"building": "A楼", "success": True, "capacity_status": "success"}]},
    ))
    with pytest.raises(RuntimeError, match="cancelled during overlay"):
        task_handlers.handle_handover_review_regenerate({}, {
            "building": "A楼", "session_id": "manual-session", "duty_date": "2026-10-06", "duty_shift": "day",
            "data_file": str(source), "capacity_source_file": str(source),
        }, lambda _: None, runtime=SimpleNamespace(raise_if_cancelled=check_cancelled))
    review_service.mark_manual_regenerated.assert_not_called()
