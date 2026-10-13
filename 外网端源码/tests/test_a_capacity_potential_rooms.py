from pathlib import Path

import pytest
from openpyxl import load_workbook

from handover_log_module.core.models import RawRow
from handover_log_module.repository.excel_reader import load_workbook_quietly
from handover_log_module.service import capacity_report_a
from handover_log_module.service.capacity_report_common import build_capacity_template_snapshot
from handover_log_module.service.capacity_room_inputs_service import CapacityRoomInputsService
from handover_log_module.service.handover_capacity_report_service import (
    HandoverCapacityReportService,
    _write_cells_with_merged_support,
)
from handover_log_module.service.review_document_state_service import ReviewDocumentStateService


@pytest.fixture
def a_template():
    path = Path(__file__).resolve().parents[1] / "A楼交接班容量报表空模板.xlsx"
    workbook = load_workbook_quietly(path)
    try:
        yield workbook, workbook["本班组"]
    finally:
        workbook.close()


def _row(c_text, value, *, room="203", d_name="进线_功率"):
    return RawRow(
        row_index=595,
        b_text=f"南通阿里保税A区A楼/A楼/挖潜机房 A-{room}",
        c_text=c_text,
        d_name=d_name,
        e_raw=str(value),
        value=value,
    )


@pytest.mark.parametrize("building,name", [
    ("A楼", "A楼交接班容量报表空模板.xlsx"),
    ("B楼", "其他楼交接班容量报表空模板.xlsx"),
    ("C楼", "其他楼交接班容量报表空模板.xlsx"),
    ("D楼", "其他楼交接班容量报表空模板.xlsx"),
    ("E楼", "E楼交接班容量报表空模板.xlsx"),
])
def test_building_template_selection(building, name, tmp_path):
    service = HandoverCapacityReportService({
        "_global_paths": {"runtime_state_root": str(tmp_path)},
        "capacity_report": {"template": {"source_path": "交接班容量报表空模板.xlsx"}},
    })
    assert service.resolve_template_path(building=building).name == name
    assert CapacityRoomInputsService._template_path_for_building(building).name == name


def test_a_appendix_fills_exact_line_head_and_room_a_meter(a_template, tmp_path):
    workbook, sheet = a_template
    snapshot = build_capacity_template_snapshot(sheet, "A楼")
    entries = snapshot["a_line_head_entries"]
    assert len(entries) == 20
    assert {item["row"] for item in entries} == set(range(159, 179))
    rows = [
        _row("A-203-A列-DC0010", 999),
        _row("A-203-A列-DC001", 999, d_name="进线_电流"),
    ]
    for i, entry in enumerate(entries):
        rows.append(_row(entry["identifier"], i * 1.25, room=entry["identifier"].split("-")[1]))
    for room, device, value in (("203", "251", 0.11), ("303", "351", 0.09), ("403", "451", 0.0)):
        rows.extend([
            _row(f"A-{device}-CRAHB-A_电量仪", 99, room=room + "0", d_name="总有功功率"),
            _row(f"A-{device}-CRAHB-B_电量仪", 99, room=room, d_name="总有功功率"),
            _row(f"A-{device}-CRAHB-A_电量仪", 99, room=room, d_name="电流"),
            _row(f"A-{device}-CRAHB-A_电量仪", value, room=room, d_name="总有功功率"),
        ])
    messages = []
    values = capacity_report_a.build_capacity_cells({
        "building": "A楼", "capacity_rows": rows,
        "template_snapshot": snapshot, "emit_log": messages.append,
    })
    for i, entry in enumerate(entries):
        assert float(values[f"E{entry['row']}"]) == i * 1.25
    assert [values[cell] for cell in ("N159", "N167", "N175")] == ["0.11", "0.09", "0"]
    assert not any("[挖潜机房]" in message for message in messages)

    protected = [f"{col}{row}" for row in range(157, 179) for col in ("A", "B", "G", "I", "L")]
    original_values = {cell: sheet[cell].value for cell in protected}
    original_styles = {cell.coordinate: cell.style_id for row in sheet.iter_rows(min_row=157, max_row=178, max_col=14) for cell in row}
    merges = set(map(str, sheet.merged_cells.ranges))
    counts = CapacityRoomInputsService.cells_from_payload({"rows": [
        {"room": "203", "powered_cabinets": "9"},
        {"room": "303", "powered_cabinets": "11"},
        {"room": "403", "powered_cabinets": "13"},
    ]}, building="A楼")
    values.update(HandoverCapacityReportService._capacity_room_overlay_values(building="A楼", handover_cells=counts))
    _write_cells_with_merged_support(sheet, values)
    path = tmp_path / "capacity.xlsx"
    workbook.save(path)
    checked = load_workbook(path)
    try:
        result = checked["本班组"]
        assert [result[cell].value for cell in ("J159", "J167", "J175")] == ["9", "11", "13"]
        assert {cell: result[cell].value for cell in protected} == original_values
        assert {cell: result[cell].style_id for cell in original_styles} == original_styles
        assert set(map(str, result.merged_cells.ranges)) == merges
        assert str(result.print_area) == str(sheet.print_area)
    finally:
        checked.close()


def test_missing_a_appendix_metrics_are_blank_not_fabricated_zero(a_template):
    _, sheet = a_template
    messages = []
    values = capacity_report_a.build_capacity_cells({
        "building": "A楼",
        "template_snapshot": build_capacity_template_snapshot(sheet, "A楼"),
        "capacity_rows": [_row("A-203-A列-DC001", 0.0), _row("A-203-A列-DC002", float("nan"))],
        "emit_log": messages.append,
    })
    assert values["E159"] == "0"
    assert values["E160"] == values["E161"] == values["N159"] == ""
    assert len([message for message in messages if "[挖潜机房]" in message]) == 1


def test_a_legacy_context_without_template_snapshot_stays_compatible():
    values = capacity_report_a.build_capacity_cells({"building": "A楼", "template_snapshot": None})
    assert "E159" not in values and "N159" not in values


def test_a_missing_metrics_log_failure_does_not_abort_generation(a_template):
    _, sheet = a_template

    def failed_logger(_message):
        raise OSError(22, "unavailable output")

    values = capacity_report_a.build_capacity_cells({
        "building": "A楼", "template_snapshot": build_capacity_template_snapshot(sheet, "A楼"),
        "capacity_rows": [], "emit_log": failed_logger,
    })
    assert values["E159"] == values["N159"] == ""


def test_new_room_counts_persist_and_participate_in_capacity_sync(tmp_path):
    config = {"_global_paths": {"runtime_state_root": str(tmp_path)}}
    document = {"capacity_room_inputs": {"rows": [
        {"room": "M1", "total_cabinets": "100", "powered_cabinets": "80", "aircon_started": "2"},
        {"room": "203", "powered_cabinets": "7"},
        {"room": "303", "powered_cabinets": "8"},
        {"room": "403", "powered_cabinets": "9"},
    ]}}
    ReviewDocumentStateService(config).persist_defaults_from_document(
        building="A楼", document=document, dirty_regions={"capacity_room_inputs": True},
    )
    restored = ReviewDocumentStateService(config).attach_capacity_room_inputs(
        document={}, session={"building": "A楼", "duty_date": "2026-10-14", "duty_shift": "day"},
    )
    cells = HandoverCapacityReportService.extract_tracked_cells_from_review_document(restored)
    assert [cells[key] for key in ("J159", "J167", "J175")] == ["7", "8", "9"]
    assert cells["AA69"] == "80"
    overlay = HandoverCapacityReportService._capacity_room_overlay_values(building="A楼", handover_cells=cells)
    assert overlay["J159"] == "7"
    assert "N159" not in overlay and "I159" not in overlay and "" not in overlay
    assert "J159=7" in HandoverCapacityReportService.capacity_input_signature(cells)


def test_legacy_room_payload_adds_a_rooms_without_overwriting_existing_counts():
    service = CapacityRoomInputsService()
    legacy = {"rows": [{"room": "M1", "total_cabinets": "100", "powered_cabinets": "80", "aircon_started": "2"}]}
    normalized = service.normalize_payload(legacy, building="A楼")
    assert len(normalized["rows"]) == 9
    assert normalized["rows"][0]["powered_cabinets"] == "80"
    assert [row["label"] for row in normalized["rows"][6:]] == ["挖潜机房203", "挖潜机房303", "挖潜机房403"]
    assert "J159" not in service.cells_from_payload(legacy, building="A楼")
    assert "" not in service.cells_from_payload(normalized, building="A楼")
    for building in ("B楼", "C楼", "D楼", "E楼"):
        assert len(service.normalize_payload(legacy, building=building)["rows"]) == 6
        assert not any(key.startswith("J") for key in service.cells_from_payload(normalized, building=building))


def test_a_powered_count_defaults_do_not_change_b_building_defaults(tmp_path):
    config = {"_global_paths": {"runtime_state_root": str(tmp_path)}}
    service = ReviewDocumentStateService(config)
    for building, rows in (
        ("B楼", [{"room": "M1", "powered_cabinets": "20"}]),
        ("A楼", [{"room": "203", "powered_cabinets": "7"}]),
    ):
        service.persist_defaults_from_document(
            building=building, document={"capacity_room_inputs": {"rows": rows}},
            dirty_regions={"capacity_room_inputs": True},
        )
    restored = ReviewDocumentStateService(config).attach_capacity_room_inputs(document={}, session={"building": "B楼"})
    cells = CapacityRoomInputsService.extract_cells_from_document(restored, building="B楼")
    assert cells["AA69"] == "20"
    assert not any(cell.startswith("J") for cell in cells)
