from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from handover_log_module.core.oil_tank_calibration import (
    B_OIL_CALIBRATION_VERSION,
    _B_TANK_LITRES,
    b_oil_litres_from_level_cm,
)
from handover_log_module.service.capacity_report_common import build_common_capacity_cell_values
from handover_log_module.service.handover_capacity_report_service import (
    HandoverCapacityReportService,
    _write_cells_with_merged_support,
)


@pytest.mark.parametrize("level_cm,litres", [
    (0, "0"), (0.5, "10"), (1.1, "10"), (2.2, "40"),
    (55, "5190"), (56.1, "5340"), (110, "13560"), (111.1, "13730"),
    (165, "21920"), (166.1, "22070"), (220, "27120"),
    (111.73, "13910"), (102.97, "12500"),
    ("111.65", "13730"), ("111.6501", "13910"), ("111.6499", "13730"),
])
def test_nearest_liquid_level_uses_centimetres_and_returns_litres(level_cm, litres):
    assert b_oil_litres_from_level_cm(level_cm) == litres


def test_calibration_has_all_200_monotonic_points():
    assert len(_B_TANK_LITRES) == 200
    assert all(a < b for a, b in zip(_B_TANK_LITRES, _B_TANK_LITRES[1:]))
    for index, litres in enumerate(_B_TANK_LITRES, 1):
        assert b_oil_litres_from_level_cm(Decimal(index) * Decimal("1.1")) == str(litres)


@pytest.mark.parametrize("value", [None, "", "bad", True, False, "NaN", "Infinity", "-Infinity", -1, 220.01])
def test_invalid_levels_are_not_clamped_to_a_valid_volume(value):
    with pytest.raises(ValueError):
        b_oil_litres_from_level_cm(value)


def _row(name, value, system="燃油自控系统"):
    return SimpleNamespace(d_name=name, e_raw=value, c_text=system)


def _service():
    return HandoverCapacityReportService.__new__(HandoverCapacityReportService)


def test_b_uses_level_rows_not_the_old_volume_aliases_and_writes_oil_cells(tmp_path):
    values = _service()._extract_current_oil_display_values(
        building="B楼", rows=[
            _row("1#油罐容积", 0.84), _row("2#油罐体积", 1.68),
            _row("油罐1液位", 220, "其他系统"),
            _row("油罐1液位", 111.73), _row("油罐2液位", 102.97),
        ], emit_log=lambda _: None,
    )
    assert values == {"first": "13910", "second": "12500"}
    cells = build_common_capacity_cell_values({"building": "B楼", "oil_current": values})
    workbook = Workbook()
    workbook.active.merge_cells("U13:W13")
    workbook.active.merge_cells("X13:Z13")
    _write_cells_with_merged_support(workbook.active, cells)
    path = tmp_path / "capacity.xlsx"
    workbook.save(path)
    workbook.close()
    result = load_workbook(path, read_only=True, data_only=True)
    try:
        assert float(result.active["U13"].value) == 13910
        assert float(result.active["X13"].value) == 12500
    finally:
        result.close()


@pytest.mark.parametrize("building,first,second", [
    ("A楼", "1000", "2000"), ("C楼", "0.84", "1.68"),
    ("D楼", "0.84", "1.68"), ("E楼", "0.84", "1.68"),
])
def test_other_buildings_keep_their_existing_oil_rules(building, first, second):
    values = _service()._extract_current_oil_display_values(
        building=building, rows=[_row("1#油罐体积", 0.84), _row("2#油罐体积", 1.68)],
        emit_log=lambda _: None,
    )
    assert values == {"first": first, "second": second}


@pytest.mark.parametrize("source,expected", [(0, "0"), (1, "1190.48"), ("0.84", "1000")])
def test_a_original_density_conversion_is_unchanged(source, expected):
    assert _service()._scale_a_oil_value(source) == expected


def test_one_invalid_b_tank_does_not_remove_the_other_tank():
    logs = []
    values = _service()._extract_current_oil_display_values(
        building="B楼", rows=[_row("油罐1液位", 221), _row("油罐2液位", 102.97)], emit_log=logs.append,
    )
    assert values == {"first": "", "second": "12500"}
    assert any("1#油罐留空" in log for log in logs)


def test_missing_level_does_not_fall_back_to_volume_metric():
    values = _service()._extract_current_oil_display_values(
        building="B楼", rows=[_row("1#油罐容积", 111.73)], emit_log=lambda _: None,
    )
    assert values == {"first": "", "second": ""}


@pytest.mark.parametrize("marked", [True, False])
@pytest.mark.parametrize("building,version", [("A楼", ""), ("B楼", B_OIL_CALIBRATION_VERSION)])
def test_previous_workbook_requires_matching_calculation_version(tmp_path, marked, building, version):
    path = tmp_path / "previous.xlsx"
    workbook = Workbook()
    workbook.active["U13"] = 13730
    workbook.active["X13"] = 12330
    workbook.properties.keywords = "existing;" + (version if marked else "")
    workbook.save(path)
    workbook.close()
    service = _service()
    service._oil_cache_service = SimpleNamespace(load_previous_values=lambda **_: {})
    service._review_session_service = SimpleNamespace(
        get_latest_session_for_context=lambda **_: {"capacity_output_file": str(path)},
    )
    service._template_cfg = lambda: {}
    current = {"first": "13910", "second": "12500"}
    previous, warning = service._load_previous_capacity_display_oil_values(
        building=building, duty_date="2026-10-23", duty_shift="day", current_display_values=current, emit_log=lambda _: None,
    )
    reusable = building != "B楼" or marked
    assert previous == ({"first": "13730", "second": "12330"} if reusable else current)
    assert bool(warning) is not reusable
